from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/os/Documents/ChatGPT/velmorax")
OUT = ROOT / "outputs" / "01a06dff-d0de-74e3-9cd7-d01803d9682b"
OUTPUT = OUT / "Version_2_Manual_Operativo_Continuidad_Velmorax.docx"
NAVY, BLUE, PALE, GRAY, BORDER = "17365D", "4472C4", "EAF2F8", "F3F4F6", "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd"); props.append(node)
    node.set(qn("w:fill"), color)


def border(cell):
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders"); props.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single"); node.set(qn("w:sz"), "6"); node.set(qn("w:color"), BORDER)
        borders.append(node)


def margins(cell):
    props = cell._tc.get_or_add_tcPr()
    tc_mar = props.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar"); props.append(tc_mar)
    for name, value in (("top", 100), ("start", 110), ("bottom", 100), ("end", 110)):
        node = OxmlElement(f"w:{name}"); node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa"); tc_mar.append(node)


def repeat_header(row):
    props = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader"); node.set(qn("w:val"), "true"); props.append(node)


def make_table(doc, headers, rows, widths, size=8.5):
    table = doc.add_table(rows=1, cols=len(headers)); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    for index, value in enumerate(headers): table.cell(0, index).text = value
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values): cells[index].text = str(value)
    repeat_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            border(cell); margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.width = Inches(widths[col_index]); shade(cell, NAVY if row_index == 0 else (PALE if row_index % 2 == 0 else "FFFFFF"))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0); paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Aptos"; run.font.size = Pt(size); run.font.color.rgb = RGBColor(255,255,255) if row_index == 0 else BLACK
                    run.bold = row_index == 0
    doc.add_paragraph()
    return table


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    props = p._p.get_or_add_pPr(); props.append(OxmlElement("w:keepNext"))
    return p


def numbered(doc, number, text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.3); p.paragraph_format.first_line_indent = Inches(-.3)
    p.add_run(f"{number}.  ").bold = True; p.add_run(text)


def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def build():
    doc = Document(); section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(.68); section.bottom_margin = Inches(.65); section.left_margin = Inches(.72); section.right_margin = Inches(.72)
    normal = doc.styles["Normal"]; normal.font.name = "Aptos"; normal.font.size = Pt(10.5); normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 23), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = doc.styles[name]; style.font.name = "Aptos Display"; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = BLACK
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Manual operativo de continuidad de Velmorax   Ficha 3229209   ")
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    for run in footer.runs: run.font.name = "Aptos"; run.font.size = Pt(8); run.font.color.rgb = RGBColor(100,100,100)

    # A deliberately different cover and information architecture.
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(50)
    r = p.add_run("ANÁLISIS Y DESARROLLO DE SOFTWARE"); r.bold = True; r.font.size = Pt(13)
    title = doc.add_paragraph(style="Title"); title.alignment = WD_ALIGN_PARAGRAPH.CENTER; title.paragraph_format.space_before = Pt(55)
    title.add_run("Manual operativo para proteger y recuperar el servicio Velmorax")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Gestión de accesos respaldos contingencia y reversa").bold = True
    doc.add_paragraph()
    info = doc.add_table(rows=5, cols=1); info.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = [
        "Equipo  Maicol Andrés Ospina  Juan David Bastidas  Emanuel Castaño",
        "Ficha  3229209",
        "Sistema seleccionado  Velmorax para operación veterinaria",
        "Tipo de documento  Manual operativo",
        "Fecha  Septiembre de 2026",
    ]
    for i, value in enumerate(values):
        info.cell(i,0).text = value; border(info.cell(i,0)); margins(info.cell(i,0)); shade(info.cell(i,0), PALE if i%2 else "FFFFFF")
        info.cell(i,0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    heading(doc, "Propósito y alcance", 1)
    doc.add_paragraph(
        "Este manual establece las decisiones que tomaríamos para mantener Velmorax disponible y recuperar la información cuando "
        "ocurra una falla. Lo planteamos para una clínica veterinaria con una o varias sedes. El documento cubre cuentas, respaldos, "
        "restauración, trabajo temporal, reversa de versiones y comunicación del incidente."
    )
    doc.add_paragraph(
        "Velmorax funciona como una aplicación web en FastAPI. En producción utiliza PostgreSQL, Caddy para HTTPS y contenedores "
        "Docker; en desarrollo puede trabajar con SQLite. La plataforma administra datos que no deberían quedar expuestos ni perderse: "
        "credenciales, mascotas y tutores, citas, historias clínicas, lotes, vencimientos, movimientos de inventario y eventos de auditoría."
    )
    heading(doc, "Mapa funcional del sistema", 2)
    make_table(doc, ["Área", "Qué permite hacer", "Quién la utiliza"], [
        ["Administración", "Crear usuarios, activar sedes, ajustar permisos, exportar auditoría y descargar respaldo.", "Owner y administrador."],
        ["Atención clínica", "Registrar pacientes, antecedentes, consultas, diagnóstico, tratamiento y evolución.", "Veterinarios y personal clínico."],
        ["Agenda", "Programar citas, asignar profesional y controlar disponibilidad.", "Recepción, personal clínico y administración."],
        ["Inventario", "Recibir lotes, controlar vencimiento y temperatura, registrar consumo, merma, traslado y abastecimiento.", "Almacén, veterinarios y administración."],
        ["Operación", "Comprobar disponibilidad, revisar logs, hacer backup y recuperar una versión estable.", "Equipo técnico."],
    ], [1.55, 3.8, 1.85], 8.6)
    doc.add_paragraph(
        "La historia clínica y la identidad del tutor tienen prioridad de protección por su carácter sensible. En inventario, el dato "
        "crítico no es solo la cantidad: también importan el lote, el vencimiento, la sede, el responsable y el movimiento anterior y posterior."
    )

    heading(doc, "Control de identidades", 1)
    heading(doc, "Matriz de acceso", 2)
    make_table(doc, ["Rol", "Cuenta a cargo de", "Puede", "No debe"], [
        ["Owner", "Propietario o director", "Administrar la red completa, aprobar inventario y consultar todos los módulos.", "Compartir su cuenta o usarla para tareas rutinarias."],
        ["Administrador", "Administrador de la clínica", "Gestionar usuarios, agenda, clínica, inventario, sedes y reportes.", "Conservar accesos de personas retiradas."],
        ["Clinical", "Coordinador veterinario", "Atender pacientes, manejar agenda, consultar y consumir inventario.", "Administrar la red o aprobar abastecimiento propio."],
        ["Inventory", "Jefe de almacén", "Gestionar lotes, movimientos, conteos y solicitudes.", "Leer historia clínica o cambiar usuarios."],
        ["Backup", "Administrador técnico", "Ejecutar copia y restauración sobre infraestructura aislada.", "Iniciar sesión como usuario clínico."],
    ], [1.05, 1.55, 3.0, 1.6], 8.1)

    heading(doc, "Ciclo de vida de una cuenta", 2)
    make_table(doc, ["Evento", "Solicitud y aprobación", "Ejecución", "Evidencia"], [
        ["Ingreso", "El jefe del área informa función, sede y módulos. Administración aprueba el rol.", "Se crea cuenta individual, activa y con cambio obligatorio de contraseña.", "Solicitud, creador, fecha, rol y sede."],
        ["Traslado o cambio", "El jefe explica el motivo. El owner aprueba si aumenta privilegios.", "Se ajustan rol, sede o permisos y se revisan tareas pendientes.", "Valor anterior, nuevo valor, actor y motivo."],
        ["Retiro", "Talento humano o jefe informa el retiro el mismo día.", "Se desactiva la cuenta, se revocan sesiones y se eliminan excepciones.", "Hora de baja, ejecutor y resultado de la prueba de acceso."],
        ["Revisión", "Administración compara mensualmente cuentas y personal vigente.", "Se suspenden cuentas sin responsable o sin uso justificado.", "Acta de revisión y acciones cerradas."],
    ], [1.05, 2.25, 2.65, 1.25], 8.1)

    heading(doc, "Acceso temporal", 2)
    doc.add_paragraph(
        "Un permiso temporal se solicitará en un registro independiente. Debe indicar usuario, permiso exacto, motivo, sede, inicio y "
        "vencimiento. El administrador podrá aprobar un apoyo operativo de bajo riesgo; el owner aprobará acceso administrativo o la "
        "facultad de aprobar inventario. El plazo será de ocho horas para soporte y máximo siete días para reemplazo de personal."
    )
    doc.add_paragraph(
        "Velmorax permite permisos personalizados por usuario, pero el vencimiento todavía requiere control operativo. Por eso el "
        "administrador programará la revocación desde el momento de la activación. El log debe dejar actor, afectado, permiso concedido, "
        "justificación, fecha inicial, fecha final y confirmación de retiro. Una ampliación se tramita como solicitud nueva."
    )

    heading(doc, "Política de copias", 1)
    doc.add_paragraph(
        "Para la base de producción elegimos backup completo en formato PostgreSQL custom. El volumen actual permite ejecutar una copia "
        "cada seis horas y una restauración completa es más sencilla que reconstruir una cadena de incrementales. El ciclo implementado "
        "cifra el archivo, genera una suma SHA 256 y realiza una restauración de verificación en una base temporal."
    )
    make_table(doc, ["Activo", "Cuándo se copia", "Conservación", "Destino"], [
        ["PostgreSQL", "Cada 6 horas y justo antes de desplegar.", "30 días; una copia mensual durante 6 meses.", "Disco separado cifrado y almacenamiento externo versionado."],
        ["SQLite de desarrollo", "Antes de migraciones o demostraciones importantes.", "10 versiones recientes.", "Carpeta de trabajo protegida fuera del repositorio."],
        ["Código y scripts", "En cada cambio aprobado y versión liberada.", "Historial permanente.", "Repositorio Git y copia remota."],
        ["Configuración", "Cuando cambie un parámetro de producción.", "Vigente y dos anteriores.", "Repositorio sin secretos y gestor seguro para claves."],
        ["Imagen desplegada", "En cada salida a producción.", "6 meses como mínimo.", "Registro de imágenes o etiqueta asociada al commit."],
        ["Manual y anexos", "En cada entrega o cambio del procedimiento.", "Historial permanente.", "Repositorio y copia documental externa."],
    ], [1.45, 2.0, 1.65, 2.1], 8.2)
    heading(doc, "Comprobación y custodia", 2)
    bullet(doc, "Tres copias: producción, respaldo local separado y copia externa.")
    bullet(doc, "Dos medios: volumen del servidor y almacenamiento remoto corporativo.")
    bullet(doc, "Una copia fuera del sitio mediante BACKUP REMOTE.")
    bullet(doc, "Clave de cifrado guardada aparte, en un gestor de contraseñas con acceso restringido.")
    bullet(doc, "Revisión diaria del último backup; una antigüedad superior a siete horas abre un incidente.")

    heading(doc, "Guía de recuperación", 1)
    doc.add_paragraph(
        "La recuperación se hará primero en un entorno aislado. El técnico no reemplazará producción hasta comprobar integridad y recibir "
        "autorización del líder del incidente. La ventana objetivo es de cuatro horas para una restauración general."
    )
    steps = [
        "Registrar hora, síntoma, módulos afectados y última operación conocida. Congelar cambios y despliegues.",
        "Determinar si la causa es aplicación, base de datos, almacenamiento, red o credenciales.",
        "Si la base responde, crear una copia de emergencia antes de cualquier reparación.",
        "Elegir el backup anterior a la falla. Validar nombre, fecha, estado verified y suma SHA 256.",
        "Descifrarlo y restaurarlo en una base temporal con el procedimiento restore backup.",
        "Comprobar tablas esenciales y conteos de usuarios, pacientes, citas, historias, lotes y movimientos.",
        "Probar inicio de sesión, lectura clínica, agenda y una consulta de inventario sin modificar producción.",
        "Autorizar el cambio, detener la aplicación y conectar la base recuperada.",
        "Levantar los servicios y revisar health live, health ready y health.",
        "Habilitar primero a usuarios de prueba, vigilar logs quince minutos y después abrir el acceso general.",
        "Reingresar los registros temporales posteriores al backup y cerrar la bitácora con tiempos y responsables.",
    ]
    for i, item in enumerate(steps, 1): numbered(doc, i, item)

    heading(doc, "Ensayo de restauración", 2)
    doc.add_paragraph(
        "El backup automatizado ya se restaura en una base temporal durante cada ciclo. Una vez al mes haremos un ensayo adicional con "
        "validación funcional. Juan David Bastidas ejecutará la restauración; Emanuel Castaño registrará tiempos, checksum, comandos, "
        "tablas y resultado; Maicol Andrés Ospina aprobará el cierre. El ensayo no tendrá acceso a usuarios de producción."
    )
    make_table(doc, ["Validación", "Aprobado cuando"], [
        ["Integridad", "El checksum coincide y pg restore no informa errores."],
        ["Estructura", "Existen las tablas operativas definidas por el script de backup."],
        ["Datos", "Los conteos son coherentes y las muestras conservan relaciones."],
        ["Funcionalidad", "Login, agenda, historia clínica e inventario responden."],
        ["Tiempo", "La recuperación termina dentro de cuatro horas."],
        ["Registro", "El acta incluye responsables, inicio, fin, evidencias, fallas y correcciones."],
    ], [1.7, 5.5], 8.8)

    heading(doc, "Análisis de impacto", 1)
    doc.add_paragraph(
        "Los objetivos se definen por proceso. RTO fija cuánto tiempo puede permanecer interrumpido; RPO limita los datos que podrían "
        "reconstruirse. Con copias cada seis horas, el RPO de los procesos que escriben datos es de seis horas."
    )
    make_table(doc, ["Proceso", "Nivel", "RTO", "RPO", "Alternativa mientras vuelve"], [
        ["Autenticación", "P1", "1 h", "No aplica", "No se comparten cuentas; se espera recuperación segura."],
        ["Historia clínica", "P1", "2 h", "6 h", "Formato clínico controlado y posterior transcripción."],
        ["Agenda", "P2", "4 h", "6 h", "Lista temporal con hora, tutor, mascota y profesional."],
        ["Inventario clínico", "P2", "4 h", "6 h", "Registro manual de lote, cantidad, paciente y responsable."],
        ["Administración", "P3", "8 h", "24 h", "Se aplazan altas y cambios no urgentes."],
        ["Reportes", "P4", "24 h", "24 h", "Se generan después de recuperar la operación."],
    ], [1.6, .55, .6, .65, 3.8], 8.0)

    heading(doc, "Respuesta por escenarios", 1)
    make_table(doc, ["Escenario y activación", "Acción inmediata", "Medida temporal", "Responsable"], [
        ["PostgreSQL no responde durante 5 minutos", "Congelar escrituras, revisar salud, logs, disco y conexión.", "Agenda e inventario en formatos numerados.", "Juan David  técnico."],
        ["Despliegue rompe login o clínica", "Detener despliegue y conservar logs y versión fallida.", "Restringir acceso hasta volver a la versión estable.", "Juan David con aprobación de Maicol."],
        ["Credencial privilegiada comprometida", "Desactivar cuenta, revocar sesiones y rotar secretos relacionados.", "Cuenta nominativa de emergencia aprobada por el líder.", "Maicol  líder del incidente."],
        ["Backup vencido o inválido", "Ejecutar copia manual, verificar espacio, clave y destino remoto.", "Suspender cambios de alto riesgo hasta tener copia válida.", "Equipo técnico."],
        ["Servidor completo fuera de servicio", "Activar infraestructura alterna y recuperar desde la copia externa.", "Comunicar indisponibilidad y operar con formatos controlados.", "Equipo completo."],
    ], [2.0, 2.45, 1.85, .9], 7.8)
    doc.add_paragraph(
        "El canal inicial será el grupo interno de incidentes. Una caída total se escalará además por llamada. Emanuel conservará la "
        "cronología; Maicol emitirá actualizaciones cada treinta minutos; Juan David informará el avance técnico sin exponer secretos ni datos clínicos."
    )

    heading(doc, "Reversa de una versión", 1)
    doc.add_paragraph(
        "Antes de cada despliegue se anotarán el commit estable, la imagen anterior, la configuración, el resultado de pruebas y el "
        "backup verificado. Se activa rollback si el servicio no supera health ready en diez minutos, aumentan los errores críticos, "
        "falla una función clínica básica o una migración deja datos inconsistentes."
    )
    rollback = [
        "Detener la liberación y bloquear nuevos cambios.",
        "Guardar logs, commit fallido y hora del error.",
        "Volver a la imagen estable anterior sin eliminar los volúmenes.",
        "Si el esquema no es compatible, restaurar la copia previa en una base aislada.",
        "Iniciar base, aplicación y proxy; comprobar los tres endpoints de salud.",
        "Ejecutar pruebas de login, pacientes, agenda, clínica e inventario.",
        "Abrir el servicio de forma gradual y vigilarlo treinta minutos.",
        "Documentar causa, decisión, resultado y condición para un nuevo intento.",
    ]
    for i, item in enumerate(rollback, 1): numbered(doc, i, item)
    doc.add_paragraph("Juan David ejecuta la reversa y Maicol autoriza la reapertura. Emanuel registra cada decisión y evidencia.")

    heading(doc, "Organización del incidente", 1)
    make_table(doc, ["Función", "Asignación en el ejercicio", "Tarea principal"], [
        ["Dirección", "Maicol Andrés Ospina", "Declara severidad, ordena contención, aprueba restauración o rollback y cierra."],
        ["Recuperación técnica", "Juan David Bastidas", "Diagnostica, ejecuta backup o reversa, valida servicios y datos."],
        ["Registro", "Emanuel Castaño", "Mantiene la cronología, reúne logs, comandos, capturas, tiempos y acuerdos."],
        ["Comunicación", "Maicol con apoyo de Emanuel", "Informa impacto confirmado, avance, siguiente actualización y contacto."],
    ], [1.3, 2.0, 3.9], 8.5)

    heading(doc, "Formato de aviso", 2)
    doc.add_paragraph(
        "Asunto  Interrupción de Velmorax en seguimiento\n\n"
        "Fecha y hora  Detectamos una falla en componente. Afecta a usuarios sedes o módulos y el impacto confirmado es descripción. "
        "El equipo aisló el componente y ahora está acción realizada. La recuperación se estima para hora o enviaremos otra actualización "
        "a las hora. Mientras tanto, siga procedimiento temporal y evite repetir operaciones. Contacto responsable  nombre  rol  canal."
    )
    doc.add_paragraph("La comunicación no incluirá credenciales, datos de pacientes, información clínica ni una causa que todavía no esté confirmada.")

    heading(doc, "Registro mínimo del incidente", 2)
    make_table(doc, ["Dato", "Contenido esperado"], [
        ["Identificación", "Código, fecha, hora, persona que reporta y severidad."],
        ["Alcance", "Sedes, usuarios, módulos y datos afectados."],
        ["Cronología", "Acción, responsable, hora, resultado y siguiente decisión."],
        ["Recuperación", "Backup o versión usada, checksum, pruebas y autorización."],
        ["Cierre", "Hora, datos reconstruidos, causa, corrección y acción preventiva."],
    ], [1.55, 5.65], 8.8)

    heading(doc, "Cierre del manual", 1)
    doc.add_paragraph(
        "Tener un archivo de respaldo no garantiza una recuperación. La copia puede estar dañada, incompleta o cifrada con una clave que "
        "nadie conserva. También puede restaurarse correctamente y aun así no permitir que la aplicación funcione. Por eso tratamos la "
        "prueba de restauración como parte del backup y no como una actividad opcional."
    )
    doc.add_paragraph(
        "La capacidad de recuperar Velmorax forma parte de la seguridad porque protege la disponibilidad y ayuda a comprobar la integridad "
        "de los datos. Los permisos reducen accesos indebidos; los logs permiten reconstruir decisiones; las copias y el rollback devuelven "
        "el sistema a un estado confiable. El plan será útil solo si los responsables lo practican y actualizan después de cada incidente."
    )
    heading(doc, "Texto para el comentario de TACA Class", 2)
    doc.add_paragraph(
        "Ante una falla crítica, mi primer paso sería declarar el incidente y detener los cambios o escrituras que puedan aumentar el daño. "
        "Después confirmaría el alcance mediante los endpoints de salud y los logs antes de decidir si corresponde restaurar datos o ejecutar rollback."
    )
    heading(doc, "Anexo de evidencias del proyecto", 2)
    doc.add_paragraph("Estas ubicaciones permiten comprobar que el manual corresponde al estado actual de Velmorax:")
    make_table(doc, ["Elemento comprobable", "Archivo o ubicación"], [
        ["Roles y permisos efectivos", "app/services/auth.py"],
        ["Alta baja y cambios de usuario", "app/routers/web.py  rutas de administración"],
        ["Backup PostgreSQL cifrado", "deploy/backup-cycle.sh"],
        ["Restauración en base temporal", "deploy/restore-backup.sh"],
        ["Frecuencia retención y destino externo", "docker-compose.production.yml"],
        ["Monitoreo de disponibilidad", "health live  health ready  health"],
        ["Procedimiento de producción", "deploy/README.md"],
        ["Pruebas automatizadas del respaldo", "tests/test_backup_configuration.py y tests/test_web.py"],
    ], [2.8, 4.4], 8.7)

    for paragraph in doc.paragraphs:
        props = paragraph._p.get_or_add_pPr()
        if props.find(qn("w:keepLines")) is None: props.append(OxmlElement("w:keepLines"))
    return doc


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    doc = build(); doc.save(OUTPUT); print(OUTPUT)

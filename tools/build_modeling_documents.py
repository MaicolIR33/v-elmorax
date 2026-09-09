from __future__ import annotations

import html
import math
import textwrap
from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output" / "documents"
ASSET_DIR = ROOT / "tmp" / "modeling" / "images"
SOURCE_DIR = ROOT / "output" / "diagram_sources"
P03_SOURCE_DIR = SOURCE_DIR / "P03"
P04_SOURCE_DIR = SOURCE_DIR / "P04"

FONT_REGULAR = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"

GREEN = "2E7D32"
DARK_GREEN = "245A39"
PALE_GREEN = "EAF4EC"
PALE_BLUE = "EDF4FA"
PALE_GRAY = "F5F6F6"
LIGHT_GRAY = "D9D9D9"
TEXT = "1E2522"
MUTED = "5F6863"
WHITE = "FFFFFF"
RED = "B33A3A"
AMBER = "B77700"
BLUE = "2E6F9E"


def ensure_dirs() -> None:
    for path in (OUTPUT_DIR, ASSET_DIR, P03_SOURCE_DIR, P04_SOURCE_DIR):
        path.mkdir(parents=True, exist_ok=True)


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size=size)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, width: int) -> list[str]:
    words = str(text).split()
    if not words:
        return [""]
    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        candidate = f"{current} {word}"
        if draw.textbbox((0, 0), candidate, font=font)[2] <= width:
            current = candidate
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def draw_centered_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    font: ImageFont.FreeTypeFont,
    fill: str = "#1E2522",
    line_spacing: int = 6,
) -> None:
    x1, y1, x2, y2 = box
    lines = wrap_text(draw, text, font, max(30, x2 - x1 - 24))
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(heights) + line_spacing * (len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, height in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=font)
        x = x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=font, fill=fill)
        y += height + line_spacing


def draw_arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[float, float],
    end: tuple[float, float],
    *,
    fill: str = "#59635E",
    width: int = 4,
    dashed: bool = False,
) -> None:
    x1, y1 = start
    x2, y2 = end
    if dashed:
        distance = math.hypot(x2 - x1, y2 - y1)
        if distance:
            ux, uy = (x2 - x1) / distance, (y2 - y1) / distance
            pos = 0.0
            while pos < distance - 14:
                stop = min(pos + 18, distance - 14)
                draw.line((x1 + ux * pos, y1 + uy * pos, x1 + ux * stop, y1 + uy * stop), fill=fill, width=width)
                pos += 30
    else:
        draw.line((x1, y1, x2, y2), fill=fill, width=width)
    angle = math.atan2(y2 - y1, x2 - x1)
    head = 15
    left = (x2 - head * math.cos(angle - math.pi / 6), y2 - head * math.sin(angle - math.pi / 6))
    right = (x2 - head * math.cos(angle + math.pi / 6), y2 - head * math.sin(angle + math.pi / 6))
    draw.polygon([(x2, y2), left, right], fill=fill)


def draw_rounded_node(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    *,
    fill: str = "#FFFFFF",
    outline: str = "#2E7D32",
    font_size: int = 25,
    bold: bool = False,
) -> None:
    draw.rounded_rectangle(box, radius=18, fill=fill, outline=outline, width=4)
    draw_centered_text(draw, box, text, pil_font(font_size, bold=bold))


def draw_gateway(
    draw: ImageDraw.ImageDraw,
    center: tuple[int, int],
    label: str,
    *,
    size: int = 54,
) -> None:
    cx, cy = center
    points = [(cx, cy - size), (cx + size, cy), (cx, cy + size), (cx - size, cy)]
    draw.polygon(points, fill="#FFF6DA", outline="#B77700")
    draw.line(points + [points[0]], fill="#B77700", width=4)
    draw.line((cx - 18, cy - 18, cx + 18, cy + 18), fill="#B77700", width=5)
    draw.line((cx + 18, cy - 18, cx - 18, cy + 18), fill="#B77700", width=5)
    if label:
        label_box = (cx - 110, cy + size + 6, cx + 110, cy + size + 68)
        draw_centered_text(draw, label_box, label, pil_font(20, bold=True), fill="#5C4300", line_spacing=3)


def draw_event(
    draw: ImageDraw.ImageDraw,
    center: tuple[int, int],
    label: str,
    *,
    end: bool = False,
) -> None:
    cx, cy = center
    radius = 30
    draw.ellipse((cx - radius, cy - radius, cx + radius, cy + radius), fill="#FFFFFF", outline="#245A39", width=7 if end else 4)
    draw_centered_text(draw, (cx - 85, cy + 35, cx + 85, cy + 100), label, pil_font(19, bold=True), fill="#245A39", line_spacing=2)


def save_drawio(path: Path, name: str, nodes: Sequence[tuple[str, str, int, int, int, int, str]], edges: Sequence[tuple[str, str, str]]) -> None:
    cells = [
        '<mxCell id="0"/>',
        '<mxCell id="1" parent="0"/>',
    ]
    for node_id, value, x, y, w, h, style in nodes:
        cells.append(
            f'<mxCell id="{html.escape(node_id)}" value="{html.escape(value)}" style="{style}" vertex="1" parent="1">'
            f'<mxGeometry x="{x}" y="{y}" width="{w}" height="{h}" as="geometry"/>'
            '</mxCell>'
        )
    for index, (source, target, label) in enumerate(edges, 1):
        cells.append(
            f'<mxCell id="e{index}" value="{html.escape(label)}" style="edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;jettySize=auto;html=1;endArrow=block;" edge="1" parent="1" source="{html.escape(source)}" target="{html.escape(target)}">'
            '<mxGeometry relative="1" as="geometry"/>'
            '</mxCell>'
        )
    xml = (
        '<mxfile host="app.diagrams.net" modified="2026-09-04T00:00:00.000Z" agent="Velmorax" version="24.7.17">'
        f'<diagram id="{html.escape(name.lower().replace(" ", "-"))}" name="{html.escape(name)}">'
        '<mxGraphModel dx="1200" dy="800" grid="1" gridSize="10" guides="1" tooltips="1" connect="1" arrows="1" fold="1" page="1" pageScale="1" pageWidth="1600" pageHeight="1000" math="0" shadow="0">'
        '<root>' + ''.join(cells) + '</root></mxGraphModel></diagram></mxfile>'
    )
    path.write_text(xml, encoding="utf-8")


def create_process_map() -> Path:
    path = ASSET_DIR / "P03_mapa_procesos.png"
    image = Image.new("RGB", (2200, 1250), "white")
    draw = ImageDraw.Draw(image)
    title_font = pil_font(42, bold=True)
    draw.text((90, 55), "Mapa de procesos de Velmorax en el flujo veterinario", font=title_font, fill="#1E2522")
    draw.text((90, 112), "Alcance revisado al 04/09/2026", font=pil_font(25), fill="#5F6863")

    bands = [
        (190, 360, "ESTRATÉGICO", "Administrar red veterinaria, sedes, usuarios, roles y configuración", "#EAF4EC"),
        (420, 810, "MISIONAL", "", "#EDF4FA"),
        (870, 1110, "APOYO", "", "#F5F6F6"),
    ]
    for top, bottom, label, description, fill in bands:
        draw.rounded_rectangle((90, top, 2110, bottom), radius=22, fill=fill, outline="#9DB8A5", width=3)
        draw.rounded_rectangle((90, top, 330, bottom), radius=22, fill="#245A39", outline="#245A39")
        draw_centered_text(draw, (105, top + 10, 315, bottom - 10), label, pil_font(27, bold=True), fill="white")
        if description:
            draw_rounded_node(draw, (400, top + 55, 2000, bottom - 55), description, fill="#FFFFFF", outline="#2E7D32", font_size=27, bold=True)

    core = [
        (400, 515, 740, 695, "Gestionar cita\nveterinaria"),
        (810, 515, 1150, 695, "Registrar paciente\ny responsable"),
        (1220, 515, 1560, 695, "Registrar atención\nveterinaria"),
        (1630, 515, 2000, 695, "Dar seguimiento y\ncerrar la atención"),
    ]
    for x1, y1, x2, y2, label in core:
        draw_rounded_node(draw, (x1, y1, x2, y2), label, fill="#FFFFFF", outline="#2E6F9E", font_size=26, bold=True)
    for left, right in zip(core, core[1:]):
        draw_arrow(draw, (left[2] + 8, (left[1] + left[3]) / 2), (right[0] - 8, (right[1] + right[3]) / 2), fill="#2E6F9E", width=5)
    draw.text((400, 455), "Solicitud del propietario", font=pil_font(22), fill="#5F6863")
    draw.text((1630, 720), "Historia, plan, control y pago registrados", font=pil_font(22), fill="#5F6863")

    support = [
        (400, 930, 900, 1045, "Gestionar inventario, lotes, vencimientos y cadena de frío"),
        (980, 930, 1450, 1045, "Auditar operaciones y exportar reportes"),
        (1530, 930, 2000, 1045, "Respaldar datos y mantener la plataforma"),
    ]
    for x1, y1, x2, y2, label in support:
        draw_rounded_node(draw, (x1, y1, x2, y2), label, fill="#FFFFFF", outline="#5F6863", font_size=23, bold=True)

    draw.rectangle((90, 1165, 125, 1200), fill="#EAF4EC", outline="#2E7D32", width=3)
    draw.text((145, 1167), "Proceso automatizado o apoyado por Velmorax", font=pil_font(22), fill="#1E2522")
    image.save(path, quality=95)

    nodes = [
        ("prc01", "PRC-01 Administrar red veterinaria", 220, 80, 1060, 90, "rounded=1;whiteSpace=wrap;html=1;fillColor=#EAF4EC;strokeColor=#2E7D32;fontStyle=1;"),
        ("prc02", "PRC-02 Gestionar cita veterinaria", 160, 280, 270, 110, "rounded=1;whiteSpace=wrap;html=1;fillColor=#EDF4FA;strokeColor=#2E6F9E;fontStyle=1;"),
        ("prc03", "PRC-03 Registrar paciente y responsable", 500, 280, 270, 110, "rounded=1;whiteSpace=wrap;html=1;fillColor=#EDF4FA;strokeColor=#2E6F9E;fontStyle=1;"),
        ("prc04", "PRC-04 Registrar atención veterinaria", 840, 280, 270, 110, "rounded=1;whiteSpace=wrap;html=1;fillColor=#EDF4FA;strokeColor=#2E6F9E;fontStyle=1;"),
        ("follow", "Seguimiento y cierre", 1180, 280, 220, 110, "rounded=1;whiteSpace=wrap;html=1;fillColor=#EDF4FA;strokeColor=#2E6F9E;fontStyle=1;"),
        ("prc05", "PRC-05 Gestionar inventario veterinario", 280, 520, 420, 110, "rounded=1;whiteSpace=wrap;html=1;fillColor=#F5F6F6;strokeColor=#5F6863;fontStyle=1;"),
        ("prc06", "PRC-06 Auditar, reportar y respaldar", 800, 520, 420, 110, "rounded=1;whiteSpace=wrap;html=1;fillColor=#F5F6F6;strokeColor=#5F6863;fontStyle=1;"),
    ]
    edges = [("prc02", "prc03", ""), ("prc03", "prc04", ""), ("prc04", "follow", "")]
    save_drawio(P03_SOURCE_DIR / "P03_A3_Mapa_de_procesos_Velmorax.drawio", "Mapa de procesos", nodes, edges)
    return path


def create_to_be_flow() -> Path:
    path = ASSET_DIR / "P03_flujo_to_be.png"
    image = Image.new("RGB", (2200, 1160), "white")
    draw = ImageDraw.Draw(image)
    draw.text((85, 45), "Flujo propuesto con Velmorax", font=pil_font(42, bold=True), fill="#1E2522")
    draw.text((85, 100), "Atención veterinaria integrada desde la cita hasta el seguimiento", font=pil_font(25), fill="#5F6863")
    lanes = [(190, 470, "PROPIETARIO"), (470, 760, "EQUIPO CLÍNICO"), (760, 1050, "VELMORAX")]
    for top, bottom, label in lanes:
        draw.rectangle((85, top, 2115, bottom), fill="#FFFFFF", outline="#A8B8AF", width=3)
        draw.rectangle((85, top, 320, bottom), fill="#EAF4EC", outline="#A8B8AF", width=3)
        draw_centered_text(draw, (100, top + 10, 305, bottom - 10), label, pil_font(24, bold=True), fill="#245A39")

    nodes = [
        ("Solicitar cita", 0, 380),
        ("Registrar cita", 1, 650),
        ("Guardar y mostrar agenda", 2, 930),
        ("Confirmar llegada", 1, 1210),
        ("Consultar ficha", 2, 1480),
        ("Atender mascota", 1, 1750),
        ("Guardar historia y control", 2, 2000),
    ]
    boxes: dict[str, tuple[int, int, int, int]] = {}
    for label, lane, x in nodes:
        top, bottom, _ = lanes[lane]
        box = (x - 105, int((top + bottom) / 2 - 58), x + 105, int((top + bottom) / 2 + 58))
        boxes[label] = box
        draw_rounded_node(draw, box, label, fill="#FFFFFF", outline="#2E7D32" if lane == 2 else "#2E6F9E", font_size=22, bold=True)
    for first, second in zip(nodes, nodes[1:]):
        a, b = boxes[first[0]], boxes[second[0]]
        draw_arrow(draw, ((a[0] + a[2]) / 2, a[3] + 5), ((b[0] + b[2]) / 2, b[1] - 5), fill="#59635E", width=4)
    draw.text((360, 1080), "Una sola captura por sede mantiene agenda, ficha, atención, pago y seguimiento en la misma base.", font=pil_font(24), fill="#1E2522")
    image.save(path, quality=95)

    src_nodes = []
    src_edges = []
    for idx, (label, lane, x) in enumerate(nodes, 1):
        src_nodes.append((f"n{idx}", label, x - 220, 150 + lane * 190, 190, 80, "rounded=1;whiteSpace=wrap;html=1;fillColor=#FFFFFF;strokeColor=#2E7D32;fontStyle=1;"))
        if idx > 1:
            src_edges.append((f"n{idx-1}", f"n{idx}", ""))
    save_drawio(P03_SOURCE_DIR / "P03_A6_Flujo_to_be_Velmorax.drawio", "Flujo to be", src_nodes, src_edges)
    return path


def create_bpmn_diagram(
    code: str,
    title: str,
    lanes: Sequence[str],
    steps: Sequence[dict],
    messages: Sequence[tuple[str, str, str]],
) -> Path:
    path = ASSET_DIR / f"P03_{code}.png"
    width, height = 2500, 1480
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 40), f"{code}  {title}", font=pil_font(40, bold=True), fill="#1E2522")
    draw.text((70, 94), "BPMN 2.0 con dos pools, lanes, eventos, tareas y compuertas exclusivas", font=pil_font(23), fill="#5F6863")

    top = 180
    lane_height = 250
    pool_label = 135
    left_label = 220
    for i, lane in enumerate(lanes):
        y1 = top + i * lane_height
        y2 = y1 + lane_height
        fill = "#F7FBF8" if i % 2 == 0 else "#F5F8FA"
        draw.rectangle((60, y1, width - 60, y2), fill=fill, outline="#9EAAA4", width=3)
        if i == 0:
            draw.rectangle((60, y1, left_label, y2), fill="#EAF4EC", outline="#9EAAA4", width=3)
            draw_centered_text(draw, (68, y1 + 8, left_label - 8, y2 - 8), f"POOL 1\n{lane}", pil_font(18, bold=True), fill="#245A39")
        else:
            draw.rectangle((pool_label, y1, left_label, y2), fill="#EAF4EC", outline="#9EAAA4", width=3)
            draw_centered_text(draw, (pool_label + 5, y1 + 8, left_label - 5, y2 - 8), lane, pil_font(19, bold=True), fill="#245A39")
    if len(lanes) > 1:
        draw.rectangle((60, top + lane_height, pool_label, top + len(lanes) * lane_height), fill="#DDEEE1", outline="#789281", width=4)
        draw_centered_text(
            draw,
            (66, top + lane_height + 8, pool_label - 6, top + len(lanes) * lane_height - 8),
            "POOL 2\nVELMORAX",
            pil_font(17, bold=True),
            fill="#245A39",
        )
        draw.line((60, top + lane_height, width - 60, top + lane_height), fill="#789281", width=7)

    node_boxes: dict[str, tuple[int, int, int, int]] = {}
    node_centers: dict[str, tuple[int, int]] = {}
    for step in steps:
        x = int(step["x"])
        lane = int(step["lane"])
        y = top + lane * lane_height + lane_height // 2
        kind = step.get("kind", "task")
        node_id = step["id"]
        label = step["label"]
        node_centers[node_id] = (x, y)
        if kind == "start":
            draw_event(draw, (x, y), label)
            node_boxes[node_id] = (x - 30, y - 30, x + 30, y + 30)
        elif kind == "end":
            draw_event(draw, (x, y), label, end=True)
            node_boxes[node_id] = (x - 30, y - 30, x + 30, y + 30)
        elif kind == "gateway":
            draw_gateway(draw, (x, y), label)
            node_boxes[node_id] = (x - 54, y - 54, x + 54, y + 54)
        else:
            box = (x - 125, y - 62, x + 125, y + 62)
            draw_rounded_node(draw, box, label, fill="#FFFFFF", outline="#2E6F9E" if lane == 0 else "#2E7D32", font_size=21, bold=True)
            node_boxes[node_id] = box

    for step in steps:
        for target, edge_label, edge_type in step.get("next", []):
            a = node_centers[step["id"]]
            b = node_centers[target]
            ax, ay = a
            bx, by = b
            if abs(bx - ax) >= abs(by - ay):
                start = (ax + (70 if bx > ax else -70), ay)
                end = (bx - (70 if bx > ax else -70), by)
            else:
                start = (ax, ay + (70 if by > ay else -70))
                end = (bx, by - (70 if by > ay else -70))
            draw_arrow(draw, start, end, fill="#555F5A", width=4, dashed=edge_type == "message")
            if edge_label:
                mx, my = (start[0] + end[0]) / 2, (start[1] + end[1]) / 2
                draw.rounded_rectangle((mx - 75, my - 24, mx + 75, my + 24), radius=8, fill="white", outline="#D9D9D9")
                draw_centered_text(draw, (int(mx - 72), int(my - 21), int(mx + 72), int(my + 21)), edge_label, pil_font(17, bold=True), fill="#5F6863")

    for source, target, label in messages:
        a, b = node_centers[source], node_centers[target]
        draw_arrow(draw, (a[0], a[1] + 68), (b[0], b[1] - 68), fill="#2E6F9E", width=4, dashed=True)
        mx, my = (a[0] + b[0]) / 2, (a[1] + b[1]) / 2
        draw_centered_text(draw, (int(mx - 110), int(my - 42), int(mx + 110), int(my + 42)), label, pil_font(18, bold=True), fill="#2E6F9E")

    legend_y = top + len(lanes) * lane_height + 35
    draw.line((100, legend_y, 210, legend_y), fill="#555F5A", width=4)
    draw.text((230, legend_y - 17), "Flujo de secuencia dentro del pool", font=pil_font(20), fill="#1E2522")
    draw_arrow(draw, (760, legend_y), (870, legend_y), fill="#2E6F9E", width=4, dashed=True)
    draw.text((895, legend_y - 17), "Flujo de mensaje entre participantes", font=pil_font(20), fill="#1E2522")
    image.save(path, quality=95)

    mx_nodes = [
        ("pool_owner", f"POOL 1 | {lanes[0]}", 10, 30, 1630, 170, "swimlane;horizontal=0;startSize=110;whiteSpace=wrap;html=1;fillColor=#F7FBF8;strokeColor=#789281;fontStyle=1;"),
        ("pool_velmorax", "POOL 2 | Velmorax", 10, 220, 1630, 350, "swimlane;horizontal=0;startSize=110;whiteSpace=wrap;html=1;fillColor=#F5F8FA;strokeColor=#789281;fontStyle=1;"),
        ("lane_business", lanes[1], 120, 220, 1520, 170, "swimlane;horizontal=0;startSize=90;whiteSpace=wrap;html=1;fillColor=#F5F8FA;strokeColor=#9EAAA4;fontStyle=1;"),
        ("lane_system", lanes[2], 120, 400, 1520, 170, "swimlane;horizontal=0;startSize=90;whiteSpace=wrap;html=1;fillColor=#F7FBF8;strokeColor=#9EAAA4;fontStyle=1;"),
    ]
    mx_edges = []
    for index, step in enumerate(steps, 1):
        style = "rounded=1;whiteSpace=wrap;html=1;fillColor=#FFFFFF;strokeColor=#2E7D32;fontStyle=1;"
        if step.get("kind") == "gateway":
            style = "rhombus;whiteSpace=wrap;html=1;fillColor=#FFF6DA;strokeColor=#B77700;fontStyle=1;"
        elif step.get("kind") in {"start", "end"}:
            style = "ellipse;whiteSpace=wrap;html=1;fillColor=#FFFFFF;strokeColor=#245A39;"
        mx_nodes.append((step["id"], step["label"], int(step["x"] / 1.5), 80 + int(step["lane"]) * 180, 150, 70, style))
        for target, edge_label, _edge_type in step.get("next", []):
            mx_edges.append((step["id"], target, edge_label))
    for source, target, label in messages:
        mx_edges.append((source, target, label))
    save_drawio(P03_SOURCE_DIR / f"P03_{code}_{title.replace(' ', '_')}.drawio", f"{code} {title}", mx_nodes, mx_edges)
    return path


def create_sequence_diagram(code: str, title: str, participants: Sequence[str], messages: Sequence[dict], fragments: Sequence[dict] | None = None) -> Path:
    path = ASSET_DIR / f"P04_{code}.png"
    width = 2400
    row_height = 92
    header_y = 185
    footer_space = 120
    height = header_y + 150 + row_height * len(messages) + footer_space
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 42), f"{code}  {title}", font=pil_font(40, bold=True), fill="#1E2522")
    draw.text((70, 98), "Diagrama de secuencia UML basado en la arquitectura actual de Velmorax", font=pil_font(23), fill="#5F6863")
    left, right = 130, width - 130
    xs = [int(left + i * (right - left) / (len(participants) - 1)) for i in range(len(participants))]
    head_font = pil_font(21, bold=True)
    for x, participant in zip(xs, participants):
        box = (x - 125, header_y, x + 125, header_y + 92)
        draw_rounded_node(draw, box, participant, fill="#EAF4EC", outline="#2E7D32", font_size=20, bold=True)
        draw.line((x, header_y + 92, x, height - 70), fill="#A5AEA9", width=3)
        for yy in range(header_y + 105, height - 70, 28):
            draw.line((x, yy, x, min(yy + 13, height - 70)), fill="#FFFFFF", width=4)

    if fragments:
        for fragment in fragments:
            start_idx, end_idx = fragment["start"], fragment["end"]
            y1 = header_y + 125 + row_height * start_idx - 22
            y2 = header_y + 125 + row_height * (end_idx + 1) - 22
            draw.rectangle((80, y1, width - 80, y2), outline="#87918C", width=3)
            label = fragment.get("label", "alt")
            draw.polygon([(80, y1), (300, y1), (270, y1 + 42), (80, y1 + 42)], fill="#F1F3F2", outline="#87918C")
            draw.text((96, y1 + 8), label, font=pil_font(20, bold=True), fill="#3F4743")

    for idx, message in enumerate(messages):
        y = header_y + 135 + idx * row_height
        src = int(message["src"])
        dst = int(message["dst"])
        x1, x2 = xs[src], xs[dst]
        is_return = message.get("type") == "Retorno"
        color = "#2E6F9E" if not is_return else "#5F6863"
        draw_arrow(draw, (x1, y), (x2, y), fill=color, width=4, dashed=is_return)
        label = f"{idx + 1}. {message['label']}"
        font = pil_font(19, bold=not is_return)
        max_width = max(250, abs(x2 - x1) - 30)
        lines = wrap_text(draw, label, font, max_width)
        text_y = y - 51
        for line in lines[:2]:
            bbox = draw.textbbox((0, 0), line, font=font)
            tx = min(x1, x2) + (abs(x2 - x1) - (bbox[2] - bbox[0])) / 2
            draw.rectangle((tx - 5, text_y - 2, tx + (bbox[2] - bbox[0]) + 5, text_y + 25), fill="white")
            draw.text((tx, text_y), line, font=font, fill=color)
            text_y += 25
    image.save(path, quality=95)

    puml = ["@startuml", f"title {code} {title}", "autonumber"]
    aliases = []
    for idx, participant in enumerate(participants):
        alias = f"P{idx+1}"
        aliases.append(alias)
        puml.append(f'participant "{participant}" as {alias}')
    fragment_starts = {frag["start"]: frag for frag in fragments or []}
    fragment_ends = {frag["end"]: frag for frag in fragments or []}
    for idx, message in enumerate(messages):
        if idx in fragment_starts:
            puml.append(f"alt {fragment_starts[idx].get('label', 'alternativa')}")
        arrow = "-->" if message.get("type") == "Retorno" else "->"
        puml.append(f"{aliases[message['src']]} {arrow} {aliases[message['dst']]}: {message['label']}")
        if idx in fragment_ends:
            puml.append("end")
    puml.append("@enduml")
    (P04_SOURCE_DIR / f"P04_{code}.puml").write_text("\n".join(puml), encoding="utf-8")
    return path


def create_activity_diagram(code: str, title: str, lanes: Sequence[str], flow: Sequence[dict]) -> Path:
    path = ASSET_DIR / f"P04_{code}.png"
    width = 2400
    top = 175
    lane_left = 110
    lane_right = width - 110
    lane_width = (lane_right - lane_left) / len(lanes)
    row_height = 180
    height = top + 110 + row_height * len(flow) + 140
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 40), f"{code}  {title}", font=pil_font(40, bold=True), fill="#1E2522")
    draw.text((70, 96), "Diagrama de actividades UML con decisiones y desenlaces explícitos", font=pil_font(23), fill="#5F6863")
    for idx, lane in enumerate(lanes):
        x1 = int(lane_left + idx * lane_width)
        x2 = int(lane_left + (idx + 1) * lane_width)
        draw.rectangle((x1, top, x2, height - 90), fill="#F9FBFA" if idx % 2 == 0 else "#F5F8FA", outline="#A4AEA8", width=3)
        draw.rectangle((x1, top, x2, top + 85), fill="#EAF4EC", outline="#A4AEA8", width=3)
        draw_centered_text(draw, (x1 + 8, top + 8, x2 - 8, top + 77), lane, pil_font(23, bold=True), fill="#245A39")

    centers: list[tuple[int, int]] = []
    for idx, item in enumerate(flow):
        lane = int(item["lane"])
        x = int(lane_left + lane * lane_width + lane_width / 2)
        y = int(top + 135 + idx * row_height)
        centers.append((x, y))
        kind = item.get("kind", "action")
        label = item["label"]
        if kind == "start":
            draw.ellipse((x - 25, y - 25, x + 25, y + 25), fill="#245A39", outline="#245A39")
            draw.text((x + 40, y - 13), label, font=pil_font(20, bold=True), fill="#245A39")
        elif kind == "end":
            draw.ellipse((x - 30, y - 30, x + 30, y + 30), fill="white", outline="#245A39", width=6)
            draw.ellipse((x - 17, y - 17, x + 17, y + 17), fill="#245A39")
            draw.text((x + 42, y - 13), label, font=pil_font(20, bold=True), fill="#245A39")
        elif kind == "decision":
            draw_gateway(draw, (x, y), label, size=50)
        else:
            draw_rounded_node(draw, (x - 205, y - 50, x + 205, y + 50), label, fill="#FFFFFF", outline="#2E6F9E" if lane == 0 else "#2E7D32", font_size=21, bold=True)
        if idx > 0:
            px, py = centers[idx - 1]
            draw_arrow(draw, (px, py + 58), (x, y - 58), fill="#59635E", width=4)
        if item.get("false"):
            branch_lane, branch_text, branch_result = item["false"]
            bx = int(lane_left + int(branch_lane) * lane_width + lane_width / 2)
            by = y + 70
            if bx == x:
                if int(branch_lane) == len(lanes) - 1:
                    bx = int(max(lane_left + 210, x - lane_width * 0.42))
                    by = y - 50
                else:
                    bx = int(min(lane_right - 210, x + lane_width * 0.38))
            branch_direction = 1 if bx > x else -1
            branch_start = (x + branch_direction * 58, y)
            branch_end = (bx - branch_direction * 205, by + 50)
            draw_arrow(draw, branch_start, branch_end, fill="#B33A3A", width=4)
            draw.text(((x + bx) / 2 - 20, y - 35), "No", font=pil_font(18, bold=True), fill="#B33A3A")
            draw_rounded_node(draw, (bx - 205, by, bx + 205, by + 100), branch_text, fill="#FFF5F5", outline="#B33A3A", font_size=19, bold=True)
            if str(branch_result).lower().startswith("fin"):
                end_y = by + 132
                draw.ellipse((bx - 18, end_y - 18, bx + 18, end_y + 18), fill="white", outline="#8F2D2D", width=4)
                draw.ellipse((bx - 9, end_y - 9, bx + 9, end_y + 9), fill="#8F2D2D")
                draw.text((bx + 28, end_y - 11), branch_result, font=pil_font(16, bold=True), fill="#8F2D2D")
            else:
                draw.text((bx - 185, by + 108), branch_result, font=pil_font(17), fill="#8F2D2D")
            if "rejoin" in item:
                target_index = int(item["rejoin"])
                target_lane = int(flow[target_index]["lane"])
                target_x = int(lane_left + target_lane * lane_width + lane_width / 2)
                target_y = int(top + 135 + target_index * row_height)
                line_y = target_y
                draw.line((bx, by + 100, bx, line_y), fill="#B33A3A", width=4)
                target_side = target_x - 205 if bx < target_x else target_x + 205
                draw_arrow(draw, (bx, line_y), (target_side, line_y), fill="#B33A3A", width=4)
        if item.get("kind") == "decision":
            draw.text((x + 64, y + 38), "Sí", font=pil_font(18, bold=True), fill="#2E7D32")
    image.save(path, quality=95)

    puml = ["@startuml", f"title {code} {title}"]

    def emit_lane(lane_index: int) -> None:
        puml.append(f'|{lanes[lane_index]}|')

    def emit_simple(item: dict) -> None:
        emit_lane(int(item["lane"]))
        kind = item.get("kind", "action")
        if kind == "start":
            puml.append("start")
        elif kind == "end":
            puml.append("stop")
        else:
            puml.append(f":{item['label']};")

    def emit_range(start_index: int, end_index: int) -> None:
        index = start_index
        while index < end_index:
            item = flow[index]
            if item.get("kind") != "decision":
                emit_simple(item)
                index += 1
                continue

            emit_lane(int(item["lane"]))
            puml.append(f"if ({item['label']}) then (Sí)")
            rejoin = item.get("rejoin")
            if rejoin is not None:
                rejoin_index = int(rejoin)
                emit_range(index + 1, rejoin_index)
                puml.append("else (No)")
                branch_lane, branch_text, branch_result = item["false"]
                emit_lane(int(branch_lane))
                puml.append(f":{branch_text};")
                if str(branch_result).lower().startswith("fin"):
                    puml.append("stop")
                puml.append("endif")
                index = rejoin_index
                continue

            if item.get("false"):
                puml.append("else (No)")
                branch_lane, branch_text, branch_result = item["false"]
                emit_lane(int(branch_lane))
                puml.append(f":{branch_text};")
                if str(branch_result).lower().startswith("fin"):
                    puml.append("stop")
            puml.append("endif")
            index += 1

    emit_range(0, len(flow))
    puml.append("@enduml")
    (P04_SOURCE_DIR / f"P04_{code}.puml").write_text("\n".join(puml), encoding="utf-8")
    return path


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = LIGHT_GRAY, size: int = 6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 85, start: int = 90, bottom: int = 85, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_paragraph_keep(paragraph, *, keep_with_next: bool = False, keep_lines: bool = True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_lines:
        p_pr.append(OxmlElement("w:keepLines"))
    if keep_with_next:
        p_pr.append(OxmlElement("w:keepNext"))


def set_run_font(run, name: str = "Arial") -> None:
    run.font.name = name
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_run_font(run)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def configure_document(doc: Document, code: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(10)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size, before, after in (("Heading 1", 16, 12, 7), ("Heading 2", 13, 10, 5), ("Heading 3", 11.5, 8, 4)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.text = f"{code}  |  Velmorax  |  Modelado del flujo veterinario"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        set_run_font(run)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "Juan Bastidas  |  04/09/2026     "
    for run in p.runs:
        set_run_font(run)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
    add_page_number(p)


def add_title_page(doc: Document, code: str, artifact_name: str, subtitle: str, destination: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("SERVICIO NACIONAL DE APRENDIZAJE SENA")
    set_run_font(r)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(DARK_GREEN)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Tecnología en Análisis y Desarrollo de Software")
    set_run_font(r)
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(MUTED)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.add_run(artifact_name)
    p_pr = p._p.get_or_add_pPr()
    p_border = p_pr.find(qn("w:pBdr"))
    if p_border is not None:
        p_pr.remove(p_border)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(subtitle)
    set_run_font(r)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(GREEN)

    metadata = [
        ["Código de la plantilla", code],
        ["Proyecto", "Velmorax"],
        ["Equipo y ficha", "Maicol Andrés Ospina y Juan David Bastidas | ficha 3229209"],
        ["Versión", "1.0"],
        ["Estado", "En revisión"],
        ["Fecha de elaboración", "04/09/2026"],
        ["Responsable", "Juan Bastidas"],
        ["Evidencia GA2 asociada", "GA2 | código específico no suministrado"],
        ["Carpeta destino", destination],
    ]
    add_table(doc, metadata, widths=[2.1, 4.7], header=False, label_column=True, font_size=9.5)
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Alcance del documento")
    set_run_font(r)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(DARK_GREEN)
    p = doc.add_paragraph(
        "El modelado se construyó a partir del comportamiento que está implementado en Velmorax al 04/09/2026. "
        "Se prioriza el acceso veterinario, la agenda, el registro de mascotas, la atención clínica, el seguimiento y el inventario de apoyo."
    )
    p.paragraph_format.space_after = Pt(6)


def add_table(
    doc: Document,
    rows: Sequence[Sequence[object]],
    *,
    widths: Sequence[float] | None = None,
    header: bool = True,
    label_column: bool = False,
    font_size: float = 8.7,
    alignments: Sequence[str] | None = None,
) -> object:
    if not rows:
        return None
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            if widths:
                cell.width = Inches(widths[c_idx])
            if header and r_idx == 0:
                set_cell_shading(cell, DARK_GREEN)
                color = WHITE
                bold = True
            elif label_column and c_idx == 0:
                set_cell_shading(cell, PALE_GREEN)
                color = TEXT
                bold = True
            else:
                set_cell_shading(cell, WHITE if r_idx % 2 else PALE_GRAY)
                color = TEXT
                bold = False
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if alignments:
                    p.alignment = {
                        "left": WD_ALIGN_PARAGRAPH.LEFT,
                        "center": WD_ALIGN_PARAGRAPH.CENTER,
                        "right": WD_ALIGN_PARAGRAPH.RIGHT,
                    }[alignments[c_idx]]
                for run in p.runs:
                    set_run_font(run)
                    run.bold = bold
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor.from_string(color)
    if header:
        set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_figure(doc: Document, image_path: Path, caption: str, width: float = 7.0) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    shape = p.add_run().add_picture(str(image_path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(caption)
    set_run_font(r)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def add_source_line(doc: Document, path: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("Archivo fuente editable  ")
    set_run_font(r)
    r.bold = True
    r.font.size = Pt(9)
    r2 = p.add_run(path)
    set_run_font(r2)
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor.from_string(MUTED)


def add_ficha(doc: Document, rows: Sequence[Sequence[str]]) -> None:
    add_table(doc, rows, widths=[2.15, 4.65], header=False, label_column=True, font_size=9.1)


def add_page_heading(doc: Document, text: str, level: int):
    heading = doc.add_heading(text, level=level)
    heading.paragraph_format.page_break_before = True
    return heading


def build_p03(diagrams: dict[str, Path]) -> Path:
    doc = Document()
    configure_document(doc, "P03")
    add_title_page(
        doc,
        "P03",
        "Mapa de procesos flujo de valor y BPMN de Velmorax",
        "Modelado del negocio aplicado al flujo veterinario",
        "documentacion/modelado/P03",
    )

    doc.add_page_break()
    doc.add_heading("Parte A Mapa de procesos y flujo de valor", level=1)
    doc.add_heading("A1 Contexto del negocio", level=2)
    add_ficha(
        doc,
        [
            ["Organización o dominio", "Clínicas y consultorios veterinarios que administran citas, pacientes, historias clínicas e insumos por sede."],
            ["Proceso principal", "Atender al paciente veterinario desde la solicitud de cita hasta el registro del tratamiento y el seguimiento."],
            ["Problema que se atiende", "PRO-01 | La información de agenda, mascota, propietario, atención e inventario puede quedar separada, repetida o sin trazabilidad por sede."],
            ["Alcance del software", "Velmorax centraliza el acceso por flujo veterinario, la agenda, el registro del paciente animal, la historia clínica, los pagos, los controles y el inventario con lotes y vencimientos."],
        ],
    )

    doc.add_heading("A2 Clasificación de procesos", level=2)
    add_table(
        doc,
        [
            ["Código", "Nivel", "Nombre del proceso", "Responsable", "Intervención del software"],
            ["PRC-01", "Estratégico", "Administrar red veterinaria, sedes, usuarios y permisos", "Owner o administrador", "Automatiza"],
            ["PRC-02", "Misional", "Gestionar cita veterinaria", "Recepción o equipo clínico", "Automatiza"],
            ["PRC-03", "Misional", "Registrar paciente y responsable", "Recepción o equipo clínico", "Automatiza"],
            ["PRC-04", "Misional", "Registrar atención veterinaria y seguimiento", "Profesional veterinario", "Automatiza"],
            ["PRC-05", "De apoyo", "Gestionar inventario veterinario", "Responsable de inventario", "Automatiza"],
            ["PRC-06", "De apoyo", "Auditar, reportar y respaldar la operación", "Administrador", "Apoya"],
        ],
        widths=[0.62, 0.8, 2.55, 1.45, 1.48],
        font_size=8.1,
        alignments=["center", "center", "left", "left", "center"],
    )

    doc.add_heading("Trazabilidad usada en los modelos", level=2)
    add_table(
        doc,
        [
            ["Código", "Historia de usuario"],
            ["HU-01", "Como integrante del equipo clínico quiero registrar y actualizar citas veterinarias por sede para organizar la atención del día."],
            ["HU-02", "Como integrante del equipo clínico quiero registrar la mascota y los datos de su responsable para disponer de una ficha veterinaria."],
            ["HU-03", "Como profesional veterinario quiero registrar la atención, el diagnóstico, el tratamiento, el peso, las vacunas, el pago y el control para conservar la trazabilidad clínica."],
            ["HU-04", "Como responsable de inventario quiero controlar insumos por lote, vencimiento, existencia y cadena de frío para evitar faltantes o uso de productos vencidos."],
            ["HU-05", "Como usuario autorizado quiero entrar por el flujo veterinario y trabajar solo con la sede y los permisos que me corresponden."],
        ],
        widths=[0.8, 6.1],
        font_size=8.7,
        alignments=["center", "left"],
    )

    add_page_heading(doc, "A3 Diagrama del mapa de procesos", level=2)
    doc.add_paragraph("El mapa separa la dirección de la red, la cadena de valor veterinaria y los procesos de apoyo. El alcance resaltado corresponde a funciones que hoy existen en Velmorax.")
    add_figure(doc, diagrams["map"], "Figura 1  Mapa de procesos del flujo veterinario de Velmorax", width=7.0)
    add_source_line(doc, "output/diagram_sources/P03/P03_A3_Mapa_de_procesos_Velmorax.drawio")

    add_page_heading(doc, "A4 Flujo de valor del proceso clave", level=2)
    add_ficha(
        doc,
        [
            ["Proceso analizado", "PRC-04 | Registrar atención veterinaria y seguimiento"],
            ["Disparador", "La mascota llega a la sede para una cita programada o una consulta prioritaria."],
            ["Resultado esperado", "El propietario recibe la orientación clínica y Velmorax conserva la atención, el pago y el control en la ficha de la mascota."],
            ["Base de los tiempos", "Estimación de modelado para comparar un flujo manual fragmentado con el flujo integrado. Debe validarse con una medición en operación."],
        ],
    )
    as_is = [
        ["N", "Actividad", "Responsable", "Procesamiento", "Espera", "Observación"],
        ["1", "Recibir la solicitud y buscar un espacio", "Recepción", "5 min", "20 min", "Consulta en agenda separada"],
        ["2", "Confirmar horario y datos básicos", "Recepción", "4 min", "10 min", "Intercambio por llamada o mensaje"],
        ["3", "Registrar la cita", "Recepción", "3 min", "0 min", "Primer registro"],
        ["4", "Volver a pedir datos de mascota y responsable", "Recepción", "6 min", "8 min", "Datos repetidos"],
        ["5", "Buscar antecedentes en notas o archivos", "Veterinario", "5 min", "12 min", "Historia dispersa"],
        ["6", "Examinar la mascota y definir el manejo", "Veterinario", "20 min", "0 min", "Actividad clínica principal"],
        ["7", "Registrar atención, pago y próximo control", "Veterinario", "8 min", "5 min", "Varios registros"],
        ["8", "Consolidar el cierre del día", "Administración", "10 min", "30 min", "Reporte manual"],
    ]
    add_table(doc, as_is, widths=[0.35, 2.3, 1.15, 0.9, 0.72, 1.48], font_size=7.8, alignments=["center", "left", "left", "center", "center", "left"])
    add_table(
        doc,
        [
            ["Medida", "Resultado estimado"],
            ["Tiempo total del ciclo", "146 minutos"],
            ["Tiempo de valor agregado", "61 minutos de procesamiento"],
            ["Eficiencia del ciclo", "61 / 146 x 100 = 41,8 %"],
        ],
        widths=[2.2, 4.7],
        font_size=9.2,
    )

    doc.add_heading("A5 Desperdicios identificados", level=2)
    add_table(
        doc,
        [
            ["Código", "Tipo", "Dónde ocurre", "Efecto medible estimado"],
            ["DES-01", "Espera", "Confirmación de cita, búsqueda de antecedentes y consolidación del cierre", "85 min de espera dentro del ciclo de referencia"],
            ["DES-02", "Doble digitación", "Agenda, ficha del paciente y nota clínica", "Hasta tres capturas de datos relacionados"],
            ["DES-03", "Transporte de información", "Paso de datos entre mensajes, notas y reportes", "Al menos cuatro cambios de medio"],
            ["DES-04", "Reproceso", "Corrección de datos incompletos o difíciles de localizar", "23 min de búsqueda y corrección en el escenario"],
        ],
        widths=[0.7, 1.2, 2.7, 2.3],
        font_size=8.2,
        alignments=["center", "center", "left", "left"],
    )

    add_page_heading(doc, "A6 Flujo propuesto y mejora esperada", level=2)
    doc.add_paragraph("El flujo propuesto corresponde a la integración que Velmorax ya ofrece. Los valores son metas de validación para una prueba operativa, no resultados medidos en producción.")
    add_figure(doc, diagrams["to_be"], "Figura 2  Flujo integrado de atención veterinaria", width=7.0)
    add_source_line(doc, "output/diagram_sources/P03/P03_A6_Flujo_to_be_Velmorax.drawio")
    add_table(
        doc,
        [
            ["Desperdicio", "Cómo lo elimina el software", "Mejora esperada y medible", "Historias"],
            ["DES-01", "Consulta agenda, ficha y cierres desde la sede seleccionada.", "Reducir la espera de 85 a 14 min, meta de 83,5 %.", "HU-01, HU-03, HU-05"],
            ["DES-02", "Reutiliza el paciente registrado al crear la atención.", "Pasar de tres capturas relacionadas a una ficha y registros asociados.", "HU-02, HU-03"],
            ["DES-03", "Mantiene agenda, historia, pagos e inventario en una sola aplicación.", "Reducir de cuatro cambios de medio a uno.", "HU-01, HU-03, HU-04"],
        ],
        widths=[0.75, 2.25, 2.6, 1.25],
        font_size=8.1,
        alignments=["center", "left", "left", "center"],
    )

    rules = {
        "BPM-01": "RN-04, RN-05",
        "BPM-02": "RN-06, RN-07",
        "BPM-03": "RN-08, RN-09",
    }
    bpmn_data = [
        (
            "BPM-01",
            "Gestionar cita veterinaria",
            [
                ["Código", "BPM-01"],
                ["Nombre del proceso", "Gestionar cita veterinaria"],
                ["Objetivo", "Registrar una cita en la sede veterinaria seleccionada y dejarla disponible en la agenda del día."],
                ["Disparador", "El propietario solicita una cita para su mascota."],
                ["Evento de fin", "La cita queda confirmada o el equipo informa el error de registro."],
                ["Participantes pools", "Propietario de la mascota | Velmorax"],
                ["Roles lanes", "Propietario | Recepción | Sistema"],
                ["Entradas", "Fecha, hora, mascota o responsable, servicio, canal, estado, especialidad, nota, sede."],
                ["Salidas", "Cita registrada, confirmación visual y evento de auditoría."],
                ["Reglas de negocio", "RN-04: el rol debe administrar agenda. RN-05: la cita veterinaria conserva organización y sede del flujo seleccionado."],
                ["Historias relacionadas", "HU-01, HU-05"],
                ["Caminos de excepción", "Datos incompletos, usuario sin permiso o error al persistir la cita."],
                ["Indicador", "Porcentaje de citas registradas sin corrección y tiempo medio de registro."],
            ],
            "output/diagram_sources/P03/P03_BPM-01_Gestionar_cita_veterinaria.drawio",
        ),
        (
            "BPM-02",
            "Registrar paciente veterinario",
            [
                ["Código", "BPM-02"],
                ["Nombre del proceso", "Registrar paciente veterinario"],
                ["Objetivo", "Crear una ficha de la mascota vinculada con su responsable y la sede veterinaria."],
                ["Disparador", "La mascota no tiene ficha disponible en la sede."],
                ["Evento de fin", "La ficha queda guardada o el equipo recibe una indicación para corregir datos."],
                ["Participantes pools", "Propietario de la mascota | Velmorax"],
                ["Roles lanes", "Propietario | Recepción o equipo clínico | Sistema"],
                ["Entradas", "Nombre, tipo Animal, especialidad Veterinaria, responsable, teléfono, especie, raza, sexo, peso y estado de vacunas."],
                ["Salidas", "Paciente disponible para seleccionar en una atención."],
                ["Reglas de negocio", "RN-06: el rol debe administrar clínica. RN-07: en el flujo veterinario se usan Animal y Veterinaria como valores del formulario."],
                ["Historias relacionadas", "HU-02, HU-05"],
                ["Caminos de excepción", "Falta el nombre obligatorio, falta la sede o el usuario no tiene permiso clínico."],
                ["Indicador", "Porcentaje de fichas completas y tiempo promedio de alta del paciente."],
            ],
            "output/diagram_sources/P03/P03_BPM-02_Registrar_paciente_veterinario.drawio",
        ),
        (
            "BPM-03",
            "Registrar atención veterinaria",
            [
                ["Código", "BPM-03"],
                ["Nombre del proceso", "Registrar atención veterinaria"],
                ["Objetivo", "Guardar la valoración, el manejo, el pago y el seguimiento de una mascota."],
                ["Disparador", "El profesional inicia la atención de una mascota en la sede."],
                ["Evento de fin", "La historia queda guardada y el propietario recibe indicaciones y control."],
                ["Participantes pools", "Propietario de la mascota | Velmorax"],
                ["Roles lanes", "Propietario | Profesional veterinario | Sistema"],
                ["Entradas", "Paciente, fecha, motivo, nota, profesional, diagnóstico, tratamiento, signos, peso, control, vacuna, pago y medio."],
                ["Salidas", "Registro clínico, última visita y peso actualizados, pago y seguimiento visibles, auditoría."],
                ["Reglas de negocio", "RN-08: la atención requiere permiso clínico y paciente existente. RN-09: si el peso actual es mayor que cero se actualiza la ficha del paciente."],
                ["Historias relacionadas", "HU-03, HU-05"],
                ["Caminos de excepción", "Paciente no localizado, datos obligatorios incompletos, permiso insuficiente o error de persistencia."],
                ["Indicador", "Porcentaje de atenciones con diagnóstico, plan y seguimiento completos."],
            ],
            "output/diagram_sources/P03/P03_BPM-03_Registrar_atención_veterinaria.drawio",
        ),
    ]

    add_page_heading(doc, "Parte B Modelado BPMN 2 0", level=1)
    doc.add_paragraph("Se modelan tres procesos críticos del flujo veterinario. Las tareas usan verbo y objeto, los mensajes entre pools se dibujan con línea punteada y cada compuerta exclusiva tiene cierre.")

    for index, (code, title, ficha, source_path) in enumerate(bpmn_data):
        if index > 0:
            add_page_heading(doc, f"Ficha del proceso {code}", level=2)
        else:
            doc.add_heading(f"Ficha del proceso {code}", level=2)
        add_ficha(doc, ficha)
        add_page_heading(doc, f"Diagrama BPMN del proceso {code}", level=2)
        add_figure(doc, diagrams[code], f"Figura {index + 3}  {code} {title}", width=7.0)
        add_source_line(doc, source_path)

    add_page_heading(doc, "B1 Verificación de la notación", level=2)
    verification = [
        ["N", "Condición", "BPM-01", "BPM-02", "BPM-03"],
        ["1", "Tiene evento de inicio y evento de fin.", "Sí", "Sí", "Sí"],
        ["2", "Las tareas usan verbo más objeto.", "Sí", "Sí", "Sí"],
        ["3", "Entre pools solo hay flujos de mensaje punteados.", "Sí", "Sí", "Sí"],
        ["4", "Cada compuerta que abre caminos tiene cierre del mismo tipo.", "Sí", "Sí", "Sí"],
        ["5", "Las salidas de compuertas exclusivas tienen condición.", "Sí", "Sí", "Sí"],
        ["6", "Se muestran caminos de excepción.", "Sí", "Sí", "Sí"],
        ["7", "Los lanes identifican al responsable.", "Sí", "Sí", "Sí"],
        ["8", "El archivo editable está guardado junto con los demás recursos.", "Sí", "Sí", "Sí"],
    ]
    add_table(doc, verification, widths=[0.35, 4.45, 0.7, 0.7, 0.7], font_size=8.2, alignments=["center", "left", "center", "center", "center"])

    doc.add_heading("B2 Registro de la validación cruzada", level=2)
    doc.add_paragraph("La validación por otro equipo no se registra como realizada porque no se suministró un equipo revisor. El artefacto permanece En revisión hasta completar esta actividad presencial.")
    add_table(
        doc,
        [
            ["Proceso", "Equipo revisor", "Qué no se entendió", "Corrección aplicada"],
            ["BPM-01", "Pendiente de asignación", "Pendiente de narración por un equipo par", "Pendiente"],
            ["BPM-02", "Pendiente de asignación", "Pendiente de narración por un equipo par", "Pendiente"],
            ["BPM-03", "Pendiente de asignación", "Pendiente de narración por un equipo par", "Pendiente"],
        ],
        widths=[0.8, 1.65, 2.45, 1.95],
        font_size=8.5,
    )

    doc.add_heading("Control de cambios", level=2)
    add_table(
        doc,
        [
            ["Versión", "Fecha", "Descripción del cambio", "Responsable"],
            ["1.0", "04/09/2026", "Versión inicial del mapa, flujo de valor y tres procesos BPMN del flujo veterinario.", "Juan Bastidas"],
        ],
        widths=[0.8, 1.2, 3.8, 1.15],
        font_size=8.7,
    )

    output = OUTPUT_DIR / "Juan_Bastidas_P03_Velmorax_Mapa_Procesos_y_BPMN.docx"
    doc.save(output)
    return output


def sequence_specs() -> list[dict]:
    common = ["Usuario", "Interfaz Jinja", "Router FastAPI", "Servicios", "Acceso a datos", "Base de datos"]
    return [
        {
            "code": "DSQ-01",
            "title": "Iniciar sesión en el flujo veterinario",
            "hu": "HU-05",
            "bpmn": "BPM-01, BPM-02 y BPM-03 como acceso previo",
            "actor": "Usuario autorizado",
            "participants": common,
            "pre": "El usuario está activo y seleccionó el flujo Veterinaria.",
            "result": "La sesión queda creada con organización, sede veterinaria y permisos; el sistema redirige al módulo permitido.",
            "errors": "Flujo no seleccionado, credenciales inválidas, bloqueo temporal, rechazo de aprobación o cambio obligatorio de clave.",
            "endpoints": "POST /login | POST /login/push-approval",
            "messages": [
                {"src": 0, "dst": 1, "label": "Seleccionar Veterinaria e ingresar credenciales", "type": "Síncrono", "data": "intent, email, password, remember_me"},
                {"src": 1, "dst": 2, "label": "POST /login", "type": "Síncrono", "data": "Formulario de acceso"},
                {"src": 2, "dst": 3, "label": "authenticate_user", "type": "Síncrono", "data": "email, password"},
                {"src": 3, "dst": 4, "label": "Consultar usuario", "type": "Síncrono", "data": "email normalizado"},
                {"src": 4, "dst": 5, "label": "SELECT users", "type": "Síncrono", "data": "estado, hash, rol, sede, intentos"},
                {"src": 5, "dst": 4, "label": "Usuario o vacío", "type": "Retorno", "data": "Fila de usuario"},
                {"src": 3, "dst": 2, "label": "Usuario autenticado o None", "type": "Retorno", "data": "Usuario y rol"},
                {"src": 2, "dst": 3, "label": "Alinear sede y calcular permisos", "type": "Síncrono", "data": "intent, organization_id, location_id, role"},
                {"src": 3, "dst": 2, "label": "Sede veterinaria y destino", "type": "Retorno", "data": "scope, permissions, path"},
                {"src": 2, "dst": 1, "label": "303 a aprobación, cambio de clave o clínica", "type": "Retorno", "data": "URL y sesión"},
            ],
            "fragments": [{"start": 6, "end": 9, "label": "alt acceso válido o rechazo"}],
        },
        {
            "code": "DSQ-02",
            "title": "Registrar cita veterinaria",
            "hu": "HU-01",
            "bpmn": "BPM-01",
            "actor": "Recepción o equipo clínico",
            "participants": common,
            "pre": "Existe una sesión activa con permiso manage_agenda y una sede veterinaria seleccionada.",
            "result": "La cita queda almacenada, se registra la auditoría y la interfaz vuelve a la agenda de la fecha elegida.",
            "errors": "Sesión vencida, rol sin permiso, datos obligatorios incompletos o error de persistencia.",
            "endpoints": "POST /appointments",
            "messages": [
                {"src": 0, "dst": 1, "label": "Completar formulario de cita", "type": "Síncrono", "data": "fecha, hora, paciente, servicio, canal, estado, nota"},
                {"src": 1, "dst": 2, "label": "POST /appointments", "type": "Síncrono", "data": "Formulario y scope"},
                {"src": 2, "dst": 3, "label": "Validar sesión y manage_agenda", "type": "Síncrono", "data": "cookie de sesión"},
                {"src": 3, "dst": 4, "label": "Consultar usuario activo", "type": "Síncrono", "data": "user_id"},
                {"src": 4, "dst": 5, "label": "SELECT users", "type": "Síncrono", "data": "usuario y rol"},
                {"src": 3, "dst": 2, "label": "Usuario y permisos", "type": "Retorno", "data": "permissions"},
                {"src": 2, "dst": 3, "label": "create_appointment", "type": "Síncrono", "data": "payload de la cita"},
                {"src": 3, "dst": 4, "label": "INSERT appointment", "type": "Síncrono", "data": "organización, sede y campos"},
                {"src": 4, "dst": 5, "label": "Ejecutar INSERT y COMMIT", "type": "Síncrono", "data": "nueva cita"},
                {"src": 2, "dst": 3, "label": "log_audit_event", "type": "Síncrono", "data": "usuario, acción, paciente, fecha"},
                {"src": 2, "dst": 1, "label": "303 /agenda con confirmación", "type": "Retorno", "data": "scope y appointment_saved=1"},
            ],
            "fragments": [{"start": 2, "end": 5, "label": "alt sesión y permiso"}],
        },
        {
            "code": "DSQ-03",
            "title": "Registrar paciente veterinario",
            "hu": "HU-02",
            "bpmn": "BPM-02",
            "actor": "Recepción o equipo clínico",
            "participants": common,
            "pre": "Existe sesión activa con permiso manage_clinical y una sede veterinaria seleccionada.",
            "result": "La mascota queda disponible en la lista de pacientes y en el selector de nuevas atenciones.",
            "errors": "Sesión vencida, rol sin permiso, nombre obligatorio ausente o error de persistencia.",
            "endpoints": "POST /patients",
            "messages": [
                {"src": 0, "dst": 1, "label": "Completar ficha de mascota", "type": "Síncrono", "data": "nombre, responsable, especie, raza, peso, vacunas"},
                {"src": 1, "dst": 2, "label": "POST /patients", "type": "Síncrono", "data": "Formulario y scope"},
                {"src": 2, "dst": 3, "label": "Validar sesión y manage_clinical", "type": "Síncrono", "data": "cookie de sesión"},
                {"src": 3, "dst": 4, "label": "Consultar usuario activo", "type": "Síncrono", "data": "user_id"},
                {"src": 4, "dst": 5, "label": "SELECT users", "type": "Síncrono", "data": "usuario y rol"},
                {"src": 3, "dst": 2, "label": "Usuario y permisos", "type": "Retorno", "data": "permissions"},
                {"src": 2, "dst": 3, "label": "create_patient", "type": "Síncrono", "data": "payload Animal y Veterinaria"},
                {"src": 3, "dst": 4, "label": "INSERT patient", "type": "Síncrono", "data": "datos de mascota, propietario y sede"},
                {"src": 4, "dst": 5, "label": "Ejecutar INSERT y COMMIT", "type": "Síncrono", "data": "nuevo paciente"},
                {"src": 2, "dst": 3, "label": "log_audit_event", "type": "Síncrono", "data": "usuario, paciente y especialidad"},
                {"src": 2, "dst": 1, "label": "303 /clinica con confirmación", "type": "Retorno", "data": "scope y patient_saved=1"},
            ],
            "fragments": [{"start": 2, "end": 5, "label": "alt sesión y permiso"}],
        },
        {
            "code": "DSQ-04",
            "title": "Registrar atención veterinaria",
            "hu": "HU-03",
            "bpmn": "BPM-03",
            "actor": "Profesional veterinario",
            "participants": common,
            "pre": "La mascota está registrada; existe sesión activa con permiso manage_clinical.",
            "result": "Se crea la historia, se actualizan última visita y peso cuando aplica, se registra auditoría y se muestra la atención.",
            "errors": "Paciente no seleccionado, sesión vencida, rol sin permiso, campos obligatorios ausentes o error de transacción.",
            "endpoints": "POST /clinical-records",
            "messages": [
                {"src": 0, "dst": 1, "label": "Completar nota clínica veterinaria", "type": "Síncrono", "data": "paciente, motivo, nota, profesional, diagnóstico, plan, peso, vacuna, pago"},
                {"src": 1, "dst": 2, "label": "POST /clinical-records", "type": "Síncrono", "data": "Formulario y scope"},
                {"src": 2, "dst": 3, "label": "Validar sesión y manage_clinical", "type": "Síncrono", "data": "cookie de sesión"},
                {"src": 3, "dst": 4, "label": "Consultar usuario activo", "type": "Síncrono", "data": "user_id"},
                {"src": 4, "dst": 5, "label": "SELECT users", "type": "Síncrono", "data": "usuario y rol"},
                {"src": 3, "dst": 2, "label": "Usuario y permisos", "type": "Retorno", "data": "permissions"},
                {"src": 2, "dst": 3, "label": "create_clinical_record", "type": "Síncrono", "data": "payload clínico"},
                {"src": 3, "dst": 4, "label": "INSERT historia y UPDATE paciente", "type": "Síncrono", "data": "atención, fecha, peso"},
                {"src": 4, "dst": 5, "label": "Ejecutar transacción y COMMIT", "type": "Síncrono", "data": "historia y ficha actualizada"},
                {"src": 2, "dst": 3, "label": "log_audit_event", "type": "Síncrono", "data": "usuario, paciente, diagnóstico y pago"},
                {"src": 2, "dst": 1, "label": "303 /clinica con confirmación", "type": "Retorno", "data": "scope, fecha y record_saved=1"},
            ],
            "fragments": [{"start": 2, "end": 5, "label": "alt sesión y permiso"}],
        },
    ]


def activity_specs() -> list[dict]:
    return [
        {
            "code": "DAC-01",
            "title": "Validar acceso al flujo veterinario",
            "hu": "HU-05",
            "bpmn": "BPM-01, BPM-02 y BPM-03",
            "sequence": "DSQ-01",
            "lanes": ["Usuario", "Router y servicios", "Datos"],
            "start": "Seleccionar el flujo Veterinaria",
            "ends": "Acceso concedido al módulo permitido | acceso rechazado | cambio de clave requerido.",
            "parallel": "No hay fork ni join. El flujo actual valida y redirige de forma síncrona.",
            "flow": [
                {"lane": 0, "kind": "start", "label": "Inicio"},
                {"lane": 0, "label": "Ingresar correo y clave"},
                {"lane": 1, "kind": "decision", "label": "¿Se seleccionó Veterinaria?", "false": (0, "Mostrar selección obligatoria", "Fin rechazado")},
                {"lane": 1, "label": "Autenticar usuario"},
                {"lane": 2, "kind": "decision", "label": "¿Usuario activo y credenciales válidas?", "false": (0, "Registrar fallo y mostrar error", "Fin rechazado")},
                {"lane": 1, "kind": "decision", "label": "¿Rol privilegiado?", "false": (1, "Continuar sin aprobación adicional", "Unión del flujo"), "rejoin": 8},
                {"lane": 0, "label": "Aprobar ingreso en segundo paso"},
                {"lane": 1, "kind": "decision", "label": "¿Aprobación concedida?", "false": (0, "Cancelar solicitud de ingreso", "Fin rechazado")},
                {"lane": 1, "label": "Alinear sede veterinaria y permisos"},
                {"lane": 1, "label": "Crear sesión y redirigir"},
                {"lane": 0, "kind": "end", "label": "Fin"},
            ],
            "decisions": [
                ["D-01", "intent debe ser veterinaria", "RN-01", "Autenticar", "Solicitar selección"],
                ["D-02", "Usuario activo, sin bloqueo y hash válido", "RN-02", "Reiniciar intentos", "Registrar fallo"],
                ["D-03", "role pertenece a admin u owner", "RN-03", "Solicitar aprobación", "Continuar"],
                ["D-04", "approve pertenece a 1, true, yes u on", "RN-03", "Crear sesión", "Rechazar ingreso"],
            ],
            "code_map": [
                ["D-01", "if", "app/routers/web.py | login_action"],
                ["D-02", "if y funciones", "app/services/auth.py | authenticate_user y register_failed_login"],
                ["D-03 y D-04", "if", "app/routers/web.py | login_action y login_push_approval_action"],
                ["Alinear sede", "función", "app/routers/web.py | align_scope_to_entry_intent"],
                ["Fork o join", "No aplica", "La implementación es síncrona"],
            ],
        },
        {
            "code": "DAC-02",
            "title": "Registrar cita veterinaria",
            "hu": "HU-01",
            "bpmn": "BPM-01",
            "sequence": "DSQ-02",
            "lanes": ["Recepción", "Router y servicios", "Datos"],
            "start": "Abrir el formulario Nueva cita en la agenda veterinaria.",
            "ends": "Cita registrada y agenda actualizada | retorno por sesión o permiso insuficiente.",
            "parallel": "No hay fork ni join. Registro, auditoría y redirección se ejecutan en orden.",
            "flow": [
                {"lane": 0, "kind": "start", "label": "Inicio"},
                {"lane": 0, "label": "Completar fecha, hora, paciente, servicio y nota"},
                {"lane": 1, "kind": "decision", "label": "¿Sesión activa?", "false": (0, "Redirigir al login", "Fin rechazado")},
                {"lane": 1, "kind": "decision", "label": "¿Tiene manage_agenda?", "false": (0, "Mostrar error de permiso", "Fin rechazado")},
                {"lane": 1, "label": "Normalizar texto y formar payload"},
                {"lane": 2, "label": "Insertar cita y confirmar transacción"},
                {"lane": 2, "label": "Registrar evento de auditoría"},
                {"lane": 1, "label": "Redirigir a agenda de la fecha"},
                {"lane": 0, "kind": "end", "label": "Fin"},
            ],
            "decisions": [
                ["D-01", "current_user_from_request devuelve usuario", "RN-02", "Validar permiso", "Ir a login"],
                ["D-02", "permissions.manage_agenda es True", "RN-04", "Guardar cita", "Mostrar error"],
            ],
            "code_map": [
                ["D-01 y D-02", "if", "app/routers/web.py | create_new_appointment"],
                ["Guardar cita", "función", "app/services/appointments.py | create_appointment"],
                ["Registrar auditoría", "función", "app/services/network.py | log_audit_event"],
                ["Ciclo de repetición", "No aplica", "No existe bucle en este flujo"],
                ["Fork o join", "No aplica", "La implementación es síncrona"],
            ],
        },
        {
            "code": "DAC-03",
            "title": "Registrar paciente veterinario",
            "hu": "HU-02",
            "bpmn": "BPM-02",
            "sequence": "DSQ-03",
            "lanes": ["Equipo clínico", "Router y servicios", "Datos"],
            "start": "Abrir Registrar paciente desde el flujo veterinario.",
            "ends": "Paciente disponible en la sede | retorno por sesión o permiso insuficiente.",
            "parallel": "No hay fork ni join. La ficha y su auditoría se guardan de forma secuencial.",
            "flow": [
                {"lane": 0, "kind": "start", "label": "Inicio"},
                {"lane": 0, "label": "Completar datos de mascota y responsable"},
                {"lane": 1, "kind": "decision", "label": "¿Sesión activa?", "false": (0, "Redirigir al login", "Fin rechazado")},
                {"lane": 1, "kind": "decision", "label": "¿Tiene manage_clinical?", "false": (0, "Mostrar error de permiso", "Fin rechazado")},
                {"lane": 1, "label": "Aplicar Animal y Veterinaria del formulario"},
                {"lane": 1, "label": "Normalizar campos y formar payload"},
                {"lane": 2, "label": "Insertar paciente y confirmar transacción"},
                {"lane": 2, "label": "Registrar evento de auditoría"},
                {"lane": 1, "label": "Redirigir a clínica"},
                {"lane": 0, "kind": "end", "label": "Fin"},
            ],
            "decisions": [
                ["D-01", "current_user_from_request devuelve usuario", "RN-02", "Validar permiso", "Ir a login"],
                ["D-02", "permissions.manage_clinical es True", "RN-06", "Guardar paciente", "Mostrar error"],
                ["D-03", "El formulario del flujo fija Animal y Veterinaria", "RN-07", "Conservar valores", "No aplica en la interfaz actual"],
            ],
            "code_map": [
                ["D-01 y D-02", "if", "app/routers/web.py | create_new_patient"],
                ["Valores veterinarios", "condición de plantilla", "app/templates/partials/clinical_section.html"],
                ["Guardar paciente", "función", "app/services/clinical.py | create_patient"],
                ["Ciclo de repetición", "No aplica", "No existe bucle en este flujo"],
                ["Fork o join", "No aplica", "La implementación es síncrona"],
            ],
        },
        {
            "code": "DAC-04",
            "title": "Registrar atención veterinaria",
            "hu": "HU-03",
            "bpmn": "BPM-03",
            "sequence": "DSQ-04",
            "lanes": ["Veterinario", "Router y servicios", "Datos"],
            "start": "Seleccionar una mascota registrada y abrir la nota clínica.",
            "ends": "Historia guardada y ficha actualizada | retorno por sesión o permiso insuficiente.",
            "parallel": "No hay fork ni join. El INSERT de historia y el UPDATE del paciente comparten una transacción; la auditoría ocurre después.",
            "flow": [
                {"lane": 0, "kind": "start", "label": "Inicio"},
                {"lane": 0, "label": "Registrar valoración, diagnóstico, plan y seguimiento"},
                {"lane": 1, "kind": "decision", "label": "¿Sesión activa?", "false": (0, "Redirigir al login", "Fin rechazado")},
                {"lane": 1, "kind": "decision", "label": "¿Tiene manage_clinical?", "false": (0, "Mostrar error de permiso", "Fin rechazado")},
                {"lane": 1, "label": "Normalizar campos y formar payload clínico"},
                {"lane": 2, "label": "Insertar registro clínico"},
                {"lane": 2, "kind": "decision", "label": "¿Peso actual mayor que cero?", "false": (1, "Conservar peso anterior", "Continuar transacción"), "rejoin": 8},
                {"lane": 2, "label": "Actualizar última visita y peso"},
                {"lane": 2, "label": "Confirmar transacción"},
                {"lane": 2, "label": "Registrar evento de auditoría"},
                {"lane": 1, "label": "Redirigir a clínica y mostrar atención"},
                {"lane": 0, "kind": "end", "label": "Fin"},
            ],
            "decisions": [
                ["D-01", "current_user_from_request devuelve usuario", "RN-02", "Validar permiso", "Ir a login"],
                ["D-02", "permissions.manage_clinical es True", "RN-08", "Guardar atención", "Mostrar error"],
                ["D-03", "current_weight_kg es mayor que 0", "RN-09", "Actualizar peso", "Conservar peso"],
            ],
            "code_map": [
                ["D-01 y D-02", "if", "app/routers/web.py | create_new_record"],
                ["Guardar atención", "función", "app/services/clinical.py | create_clinical_record"],
                ["D-03", "CASE WHEN", "app/services/clinical.py | UPDATE patients"],
                ["Transacción", "INSERT, UPDATE y commit", "app/services/clinical.py"],
                ["Fork o join", "No aplica", "La implementación es síncrona"],
            ],
        },
    ]


def build_p04(sequence_images: dict[str, Path], activity_images: dict[str, Path]) -> Path:
    doc = Document()
    configure_document(doc, "P04")
    add_title_page(
        doc,
        "P04",
        "Diagramas UML de secuencia y actividades de Velmorax",
        "Modelado del comportamiento aplicado al flujo veterinario",
        "documentacion/modelado/P04",
    )

    doc.add_page_break()
    doc.add_heading("Parte A Diagramas de secuencia", level=1)
    doc.add_heading("A1 Participantes técnicos del sistema", level=2)
    add_table(
        doc,
        [
            ["Código", "Participante", "Responsabilidad", "Capa"],
            ["CMP-01", "Interfaz Jinja en el navegador", "Mostrar formularios, enviar solicitudes HTTP y presentar respuestas.", "Presentación"],
            ["CMP-02", "Router FastAPI", "Recibir endpoints, leer sesión, validar permisos, formar payloads y redirigir.", "Aplicación"],
            ["CMP-03", "Servicios de negocio", "Autenticar, gestionar agenda, pacientes, historias, alcance y auditoría.", "Dominio"],
            ["CMP-04", "DatabaseConnection y SQL", "Normalizar consultas y ejecutar el acceso a datos usado por los servicios.", "Persistencia"],
            ["CMP-05", "SQLite o PostgreSQL", "Conservar usuarios, sedes, citas, pacientes, historias, inventario y auditoría.", "Persistencia"],
            ["EXT-01", "Correo SMTP o webhook SMS", "Enviar enlaces de recuperación cuando están configurados; no participa en los cuatro flujos principales modelados.", "Externo"],
        ],
        widths=[0.7, 1.55, 3.3, 1.25],
        font_size=8.2,
        alignments=["center", "left", "left", "center"],
    )
    doc.add_paragraph("La arquitectura real es una aplicación web renderizada en servidor. Los endpoints reciben formularios, usan sesiones y devuelven redirecciones 303; no se modela una API JSON que todavía no existe.")

    seqs = sequence_specs()
    for index, spec in enumerate(seqs):
        add_page_heading(doc, f"A2 Ficha del diagrama {spec['code']}", level=2)
        add_ficha(
            doc,
            [
                ["Código", spec["code"]],
                ["Nombre del flujo", spec["title"]],
                ["Historia de usuario", spec["hu"]],
                ["Proceso BPMN relacionado", spec["bpmn"]],
                ["Actor que inicia", spec["actor"]],
                ["Participantes", "CMP-01, CMP-02, CMP-03, CMP-04 y CMP-05"],
                ["Precondiciones", spec["pre"]],
                ["Resultado esperado", spec["result"]],
                ["Escenarios de error", spec["errors"]],
                ["Endpoints involucrados", spec["endpoints"]],
            ],
        )
        doc.add_heading(f"Diagrama de secuencia {spec['code']}", level=2)
        add_figure(doc, sequence_images[spec["code"]], f"Figura {index + 1}  {spec['code']} {spec['title']}", width=7.0)
        source_file = P04_SOURCE_DIR / f"P04_{spec['code']}.puml"
        add_source_line(doc, str(source_file.relative_to(ROOT)))

        add_page_heading(doc, f"Mensajes del diagrama {spec['code']}", level=2)
        rows = [["N", "Origen", "Destino", "Mensaje", "Tipo", "Datos que viajan"]]
        for n, message in enumerate(spec["messages"], 1):
            rows.append([
                str(n),
                spec["participants"][message["src"]],
                spec["participants"][message["dst"]],
                message["label"],
                message["type"],
                message["data"],
            ])
        add_table(doc, rows, widths=[0.35, 1.05, 1.05, 1.75, 0.75, 2.0], font_size=7.3, alignments=["center", "left", "left", "left", "center", "left"])

    add_page_heading(doc, "A3 Verificación de coherencia arquitectónica", level=2)
    add_table(
        doc,
        [
            ["N", "Condición", "Cumple", "Evidencia del modelo"],
            ["1", "Las capas corresponden con el estilo arquitectónico actual.", "Sí", "Presentación, router, servicios, acceso a datos y base de datos."],
            ["2", "No hay saltos de capa desde la interfaz a la base de datos.", "Sí", "Todos los diagramas pasan por Router y Servicios."],
            ["3", "Los mensajes asíncronos se usan solo cuando no se espera respuesta.", "Sí", "No se dibujaron mensajes asíncronos porque estos flujos son síncronos."],
            ["4", "Los errores están representados con fragmentos alt u opt.", "Sí", "Cada secuencia incluye alternativa de sesión o permiso; DSQ-01 amplía autenticación."],
            ["5", "Cada diagrama se asocia con una historia existente en este modelado.", "Sí", "HU-01, HU-02, HU-03 y HU-05."],
        ],
        widths=[0.35, 3.55, 0.65, 2.35],
        font_size=8.1,
        alignments=["center", "left", "center", "left"],
    )
    add_table(
        doc,
        [
            ["Inconsistencia detectada", "Causa", "Decisión"],
            ["No existe una clase Repository independiente.", "Los servicios ejecutan SQL mediante get_connection y DatabaseConnection.", "Corregir el diagrama para mostrar CMP-04 como adaptador de acceso real, sin inventar un repositorio."],
            ["Los formularios no consumen una API JSON.", "La interfaz es Jinja y los POST devuelven redirecciones 303.", "Corregir los mensajes para mostrar formularios HTTP y redirecciones."],
            ["La auditoría se registra después de la operación principal.", "log_audit_event es una llamada separada.", "Mostrarla como mensaje secuencial posterior; no presentarla como parte de la misma transacción."],
        ],
        widths=[2.1, 2.2, 2.7],
        font_size=8.3,
    )

    acts = activity_specs()
    add_page_heading(doc, "Parte B Diagramas de actividades", level=1)
    for index, spec in enumerate(acts):
        if index > 0:
            add_page_heading(doc, f"B1 Ficha del diagrama {spec['code']}", level=2)
        else:
            doc.add_heading(f"B1 Ficha del diagrama {spec['code']}", level=2)
        add_ficha(
            doc,
            [
                ["Código", spec["code"]],
                ["Nombre de la actividad", spec["title"]],
                ["Historia de usuario", spec["hu"]],
                ["Proceso BPMN que detalla", spec["bpmn"]],
                ["Diagrama de secuencia relacionado", spec["sequence"]],
                ["Particiones swimlanes", " | ".join(spec["lanes"])],
                ["Nodo inicial", spec["start"]],
                ["Nodos finales", spec["ends"]],
                ["Actividades paralelas", spec["parallel"]],
            ],
        )
        add_page_heading(doc, f"Diagrama de actividades {spec['code']}", level=2)
        add_figure(doc, activity_images[spec["code"]], f"Figura {index + 5}  {spec['code']} {spec['title']}", width=7.0)
        source_file = P04_SOURCE_DIR / f"P04_{spec['code']}.puml"
        add_source_line(doc, str(source_file.relative_to(ROOT)))

        add_page_heading(doc, f"Decisiones del diagrama {spec['code']}", level=2)
        add_table(
            doc,
            [["Decisión", "Condición evaluada", "Regla", "Camino verdadero", "Camino falso"]] + spec["decisions"],
            widths=[0.65, 2.35, 0.8, 1.6, 1.6],
            font_size=7.8,
            alignments=["center", "left", "center", "left", "left"],
        )
        doc.add_heading("Del diagrama al código", level=2)
        add_table(
            doc,
            [["Elemento del diagrama", "Estructura de código", "Componente o archivo"]] + spec["code_map"],
            widths=[1.55, 1.7, 3.8],
            font_size=8.1,
            alignments=["left", "center", "left"],
        )

    add_page_heading(doc, "Parte C Consolidado del modelado", level=1)
    consolidated = [["Código", "Nombre del diagrama", "Tipo", "Historia", "Archivo fuente"]]
    for spec in seqs:
        source_file = P04_SOURCE_DIR / f"P04_{spec['code']}.puml"
        consolidated.append([spec["code"], spec["title"], "Secuencia", spec["hu"], str(source_file.relative_to(ROOT))])
    for spec in acts:
        source_file = P04_SOURCE_DIR / f"P04_{spec['code']}.puml"
        consolidated.append([spec["code"], spec["title"], "Actividades", spec["hu"], str(source_file.relative_to(ROOT))])
    add_table(doc, consolidated, widths=[0.65, 2.25, 0.85, 0.75, 2.5], font_size=7.8, alignments=["center", "left", "center", "center", "left"])

    doc.add_heading("Control de cambios", level=2)
    add_table(
        doc,
        [
            ["Versión", "Fecha", "Descripción del cambio", "Responsable"],
            ["1.0", "04/09/2026", "Versión inicial con cuatro secuencias, cuatro actividades y trazabilidad al flujo veterinario actual.", "Juan Bastidas"],
        ],
        widths=[0.8, 1.2, 3.8, 1.15],
        font_size=8.7,
    )

    output = OUTPUT_DIR / "Juan_Bastidas_P04_Velmorax_UML_Secuencia_y_Actividades.docx"
    doc.save(output)
    return output


def generate_all() -> tuple[Path, Path]:
    ensure_dirs()
    diagrams: dict[str, Path] = {
        "map": create_process_map(),
        "to_be": create_to_be_flow(),
    }

    bpmn_specs = [
        (
            "BPM-01",
            "Gestionar cita veterinaria",
            ["Propietario", "Recepción", "Sistema"],
            [
                {"id": "s", "label": "Solicitud", "x": 290, "lane": 0, "kind": "start", "next": [("request", "", "sequence")]},
                {"id": "request", "label": "Solicitar cita", "x": 540, "lane": 0, "next": [("owner_end", "", "sequence")]},
                {"id": "owner_end", "label": "Respuesta", "x": 2200, "lane": 0, "kind": "end"},
                {"id": "receive", "label": "Recibir solicitud", "x": 540, "lane": 1, "next": [("check", "", "sequence")]},
                {"id": "check", "label": "¿Datos completos?", "x": 850, "lane": 1, "kind": "gateway", "next": [("merge", "Sí", "sequence"), ("correct", "No", "sequence")]},
                {"id": "correct", "label": "Solicitar corrección", "x": 850, "lane": 2, "next": [("merge", "Corregido", "sequence")]},
                {"id": "merge", "label": "Unir caminos", "x": 1180, "lane": 1, "kind": "gateway", "next": [("register", "", "sequence")]},
                {"id": "register", "label": "Registrar cita", "x": 1470, "lane": 1, "next": [("save", "", "sequence")]},
                {"id": "save", "label": "Guardar cita y auditoría", "x": 1770, "lane": 2, "next": [("confirm", "", "sequence")]},
                {"id": "confirm", "label": "Confirmar cita", "x": 2070, "lane": 1, "kind": "task", "next": [("end", "", "sequence")]},
                {"id": "end", "label": "Cita registrada", "x": 2310, "lane": 1, "kind": "end"},
            ],
            [("request", "receive", "Solicitud"), ("confirm", "owner_end", "Confirmación")],
        ),
        (
            "BPM-02",
            "Registrar paciente veterinario",
            ["Propietario", "Equipo clínico", "Sistema"],
            [
                {"id": "s", "label": "Sin ficha", "x": 280, "lane": 0, "kind": "start", "next": [("provide", "", "sequence")]},
                {"id": "provide", "label": "Entregar datos de la mascota", "x": 560, "lane": 0, "next": [("owner_end", "", "sequence")]},
                {"id": "owner_end", "label": "Ficha disponible", "x": 2200, "lane": 0, "kind": "end"},
                {"id": "capture", "label": "Capturar datos", "x": 560, "lane": 1, "next": [("check", "", "sequence")]},
                {"id": "check", "label": "¿Datos requeridos?", "x": 900, "lane": 1, "kind": "gateway", "next": [("merge", "Sí", "sequence"), ("complete", "No", "sequence")]},
                {"id": "complete", "label": "Solicitar datos faltantes", "x": 900, "lane": 2, "next": [("merge", "Completos", "sequence")]},
                {"id": "merge", "label": "Unir caminos", "x": 1240, "lane": 1, "kind": "gateway", "next": [("save", "", "sequence")]},
                {"id": "save", "label": "Guardar paciente", "x": 1580, "lane": 2, "next": [("show", "", "sequence")]},
                {"id": "show", "label": "Mostrar ficha", "x": 1930, "lane": 1, "next": [("end", "", "sequence")]},
                {"id": "end", "label": "Paciente registrado", "x": 2270, "lane": 1, "kind": "end"},
            ],
            [("provide", "capture", "Datos"), ("show", "owner_end", "Ficha")],
        ),
        (
            "BPM-03",
            "Registrar atención veterinaria",
            ["Propietario", "Veterinario", "Sistema"],
            [
                {"id": "s", "label": "Llegada", "x": 260, "lane": 0, "kind": "start", "next": [("present", "", "sequence")]},
                {"id": "present", "label": "Presentar mascota", "x": 500, "lane": 0, "next": [("owner_end", "", "sequence")]},
                {"id": "owner_end", "label": "Plan recibido", "x": 2250, "lane": 0, "kind": "end"},
                {"id": "find", "label": "Buscar paciente", "x": 500, "lane": 1, "next": [("exists", "", "sequence")]},
                {"id": "exists", "label": "¿Paciente existe?", "x": 790, "lane": 1, "kind": "gateway", "next": [("merge1", "Sí", "sequence"), ("register", "No", "sequence")]},
                {"id": "register", "label": "Registrar paciente", "x": 790, "lane": 2, "next": [("merge1", "Registrado", "sequence")]},
                {"id": "merge1", "label": "Unir caminos", "x": 1060, "lane": 1, "kind": "gateway", "next": [("examine", "", "sequence")]},
                {"id": "examine", "label": "Examinar mascota", "x": 1330, "lane": 1, "next": [("record", "", "sequence")]},
                {"id": "record", "label": "Registrar atención", "x": 1600, "lane": 1, "next": [("save", "", "sequence")]},
                {"id": "save", "label": "Guardar historia y actualizar ficha", "x": 1880, "lane": 2, "next": [("explain", "", "sequence")]},
                {"id": "explain", "label": "Explicar tratamiento y control", "x": 2140, "lane": 1, "next": [("end", "", "sequence")]},
                {"id": "end", "label": "Atención registrada", "x": 2350, "lane": 1, "kind": "end"},
            ],
            [("present", "find", "Ingreso"), ("explain", "owner_end", "Indicaciones")],
        ),
    ]
    for code, title, lanes, steps, messages in bpmn_specs:
        diagrams[code] = create_bpmn_diagram(code, title, lanes, steps, messages)

    sequence_images: dict[str, Path] = {}
    for spec in sequence_specs():
        sequence_images[spec["code"]] = create_sequence_diagram(
            spec["code"], spec["title"], spec["participants"], spec["messages"], spec["fragments"]
        )

    activity_images: dict[str, Path] = {}
    for spec in activity_specs():
        activity_images[spec["code"]] = create_activity_diagram(spec["code"], spec["title"], spec["lanes"], spec["flow"])

    p03 = build_p03(diagrams)
    p04 = build_p04(sequence_images, activity_images)
    return p03, p04


if __name__ == "__main__":
    p03_output, p04_output = generate_all()
    print(p03_output)
    print(p04_output)

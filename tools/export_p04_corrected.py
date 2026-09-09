from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "output" / "documents" / "Juan_Bastidas_P04_Velmorax_UML_Secuencia_y_Actividades.docx"
OUTPUT = ROOT / "output" / "documents" / "P04_Velmorax_CORREGIDO_Maicol_Ospina_y_Juan_David_Bastidas_3229209.docx"


document = Document(SOURCE)
updated = False
for table in document.tables:
    for row in table.rows:
        if row.cells[0].text.strip() != "Equipo y ficha":
            continue
        paragraph = row.cells[1].paragraphs[0]
        first_run = paragraph.runs[0]
        first_run.text = "Maicol Andrés Ospina y Juan David Bastidas\nFicha 3229209"
        for extra_run in paragraph.runs[1:]:
            extra_run._element.getparent().remove(extra_run._element)
        updated = True
        break
    if updated:
        break

if not updated:
    raise RuntimeError("No se encontró el campo Equipo y ficha en P04")

document.save(OUTPUT)
print(OUTPUT)

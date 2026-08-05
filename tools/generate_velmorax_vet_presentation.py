from __future__ import annotations

import sqlite3
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "expo_velmorax_veterinaria"
ASSET_DIR = OUT_DIR / "assets"
PPTX_PATH = OUT_DIR / "Velmorax_Flujo_Veterinaria_SENA.pptx"
DB_PATH = ROOT / "data" / "velmorax.db"
LOGO_PATH = ROOT / "app" / "static" / "img" / "velmorax-logo-photo-clean.png"

COLORS = {
    "ink": RGBColor(20, 31, 39),
    "muted": RGBColor(89, 103, 112),
    "teal": RGBColor(24, 165, 159),
    "teal_dark": RGBColor(11, 95, 105),
    "green": RGBColor(61, 156, 95),
    "yellow": RGBColor(227, 174, 58),
    "red": RGBColor(209, 74, 74),
    "blue": RGBColor(65, 111, 180),
    "soft": RGBColor(238, 246, 246),
    "white": RGBColor(255, 255, 255),
    "line": RGBColor(210, 224, 226),
}


def connect() -> sqlite3.Connection | None:
    if not DB_PATH.exists() or DB_PATH.stat().st_size == 0:
        return None
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    return con


def scalar(con: sqlite3.Connection | None, query: str) -> int:
    if con is None:
        return 0
    try:
        return int(con.execute(query).fetchone()[0])
    except sqlite3.Error:
        return 0


def load_metrics() -> dict:
    con = connect()
    metrics = {
        "patients": scalar(con, "SELECT COUNT(*) FROM patients"),
        "vet_patients": scalar(con, "SELECT COUNT(*) FROM patients WHERE specialty='Veterinaria'"),
        "records": scalar(con, "SELECT COUNT(*) FROM clinical_records"),
        "vet_records": scalar(con, "SELECT COUNT(*) FROM clinical_records WHERE specialty='Veterinaria'"),
        "appointments": scalar(con, "SELECT COUNT(*) FROM appointments"),
        "inventory": scalar(con, "SELECT COUNT(*) FROM inventory_items"),
        "users": scalar(con, "SELECT COUNT(*) FROM users"),
        "locations": scalar(con, "SELECT COUNT(*) FROM locations"),
        "organizations": scalar(con, "SELECT COUNT(*) FROM organizations"),
    }
    if con is not None:
        con.close()
    return metrics


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def rounded(draw: ImageDraw.ImageDraw, box, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def save_mockup_login() -> Path:
    path = ASSET_DIR / "mockup_login.png"
    img = Image.new("RGB", (1500, 900), "#eef6f6")
    d = ImageDraw.Draw(img)
    rounded(d, (70, 70, 1430, 830), 32, "#ffffff", "#d3e6e7", 2)
    d.text((120, 120), "Velmorax", fill="#0b5f69", font=font(58, True))
    d.text((120, 190), "Ingreso por flujo", fill="#596770", font=font(30))
    x = 120
    cards = [("01", "Odontología", "Dental", "#f8fafb"), ("02", "Veterinaria", "Animal", "#dff4f1"), ("03", "Consulta general", "General", "#f8fafb")]
    for idx, title, tag, fill in cards:
        rounded(d, (x, 300, x + 300, 540), 24, fill, "#18a59f" if title == "Veterinaria" else "#d3e6e7", 4 if title == "Veterinaria" else 2)
        d.text((x + 28, 328), idx, fill="#18a59f", font=font(36, True))
        d.text((x + 28, 410), title, fill="#141f27", font=font(34, True))
        d.text((x + 28, 468), tag, fill="#596770", font=font(24))
        x += 335
    rounded(d, (1120, 150, 1360, 710), 28, "#f8fbfb", "#d3e6e7", 2)
    d.text((1160, 210), "Ingreso veterinario", fill="#141f27", font=font(28, True))
    d.text((1160, 285), "Correo", fill="#596770", font=font(22))
    rounded(d, (1160, 320, 1320, 365), 12, "#ffffff", "#d3e6e7")
    d.text((1175, 330), "clinica@...", fill="#8a9aa2", font=font(18))
    d.text((1160, 400), "Clave", fill="#596770", font=font(22))
    rounded(d, (1160, 435, 1320, 480), 12, "#ffffff", "#d3e6e7")
    d.text((1175, 445), "••••••••", fill="#8a9aa2", font=font(18))
    rounded(d, (1160, 550, 1320, 615), 18, "#18a59f")
    d.text((1188, 567), "Entrar", fill="#ffffff", font=font(24, True))
    img.save(path)
    return path


def save_mockup_dashboard() -> Path:
    path = ASSET_DIR / "mockup_dashboard.png"
    img = Image.new("RGB", (1500, 900), "#f7fbfb")
    d = ImageDraw.Draw(img)
    rounded(d, (60, 55, 1440, 845), 28, "#ffffff", "#d3e6e7", 2)
    d.text((110, 100), "Centro de operaciones veterinaria", fill="#141f27", font=font(44, True))
    kpis = [("Citas de hoy", "13"), ("Pacientes visibles", "8"), ("Alertas inventario", "17"), ("Seguimientos", "3")]
    x = 110
    for label, value in kpis:
        rounded(d, (x, 200, x + 285, 360), 20, "#eef6f6", "#d3e6e7")
        d.text((x + 25, 225), label, fill="#596770", font=font(24))
        d.text((x + 25, 265), value, fill="#0b5f69", font=font(58, True))
        x += 315
    rounded(d, (110, 430, 790, 760), 22, "#ffffff", "#d3e6e7")
    d.text((145, 465), "Agenda breve", fill="#141f27", font=font(30, True))
    for i, line in enumerate(["09:20  Max / Ana Rojas - Vacunación", "11:40  Mia / David León - Control", "15:00  Rocky / Julián Toro - Revisión"]):
        y = 535 + i * 68
        d.line((145, y + 42, 745, y + 42), fill="#d3e6e7", width=2)
        d.text((145, y), line, fill="#596770", font=font(22))
    rounded(d, (840, 430, 1320, 760), 22, "#ffffff", "#d3e6e7")
    d.text((875, 465), "Ficha activa", fill="#141f27", font=font(30, True))
    d.text((875, 540), "Max", fill="#0b5f69", font=font(48, True))
    d.text((875, 605), "Canino | Labrador | 28.4 kg", fill="#596770", font=font(25))
    d.text((875, 655), "Vacunas: Al día", fill="#3d9c5f", font=font(25, True))
    img.save(path)
    return path


def save_mockup_agenda() -> Path:
    path = ASSET_DIR / "mockup_agenda.png"
    img = Image.new("RGB", (1500, 900), "#f7fbfb")
    d = ImageDraw.Draw(img)
    d.text((95, 85), "Agenda veterinaria", fill="#141f27", font=font(50, True))
    rounded(d, (90, 170, 1020, 800), 24, "#ffffff", "#d3e6e7", 2)
    entries = [
        ("08:00", "Laura Mendoza", "Control odontológico", "#416fb4"),
        ("09:20", "Max / Ana Rojas", "Vacunación veterinaria", "#18a59f"),
        ("11:40", "Mia / David León", "Revisión postoperatoria", "#18a59f"),
        ("15:00", "Rocky / Julián Toro", "Control de peso", "#e3ae3a"),
    ]
    for i, (time, patient, service, color) in enumerate(entries):
        y = 220 + i * 130
        d.text((135, y), time, fill=color, font=font(32, True))
        rounded(d, (245, y - 10, 940, y + 82), 18, "#f8fbfb", "#d3e6e7")
        d.text((280, y + 5), patient, fill="#141f27", font=font(28, True))
        d.text((280, y + 45), service, fill="#596770", font=font(22))
    rounded(d, (1060, 170, 1405, 800), 24, "#eef6f6", "#d3e6e7", 2)
    d.text((1100, 220), "Nueva cita", fill="#141f27", font=font(34, True))
    for i, label in enumerate(["Fecha", "Hora", "Mascota / responsable", "Servicio", "Estado"]):
        y = 300 + i * 82
        d.text((1100, y), label, fill="#596770", font=font(20))
        rounded(d, (1100, y + 28, 1358, y + 68), 10, "#ffffff", "#d3e6e7")
    img.save(path)
    return path


def save_mockup_clinical() -> Path:
    path = ASSET_DIR / "mockup_clinical.png"
    img = Image.new("RGB", (1500, 900), "#f7fbfb")
    d = ImageDraw.Draw(img)
    d.text((95, 80), "Ficha clínica veterinaria", fill="#141f27", font=font(50, True))
    rounded(d, (90, 170, 700, 790), 24, "#ffffff", "#d3e6e7", 2)
    d.text((130, 220), "Registrar paciente", fill="#141f27", font=font(34, True))
    fields = ["Nombre visible", "Responsable", "Especie", "Raza", "Peso kg", "Vacunas"]
    for i, label in enumerate(fields):
        y = 300 + i * 72
        d.text((130, y), label, fill="#596770", font=font(21))
        rounded(d, (330, y - 8, 645, y + 35), 10, "#f8fbfb", "#d3e6e7")
    rounded(d, (760, 170, 1410, 790), 24, "#ffffff", "#d3e6e7", 2)
    d.text((805, 220), "Atención clínica", fill="#141f27", font=font(34, True))
    tags = [("Peso actual", "28.4 kg", "#18a59f"), ("Vacunas", "Al día", "#3d9c5f"), ("Estado", "Seguimiento", "#e3ae3a")]
    x = 805
    for label, value, color in tags:
        rounded(d, (x, 300, x + 175, 405), 16, "#eef6f6", "#d3e6e7")
        d.text((x + 18, 320), label, fill="#596770", font=font(18))
        d.text((x + 18, 352), value, fill=color, font=font(24, True))
        x += 195
    d.text((805, 465), "Motivo, diagnóstico, plan, fórmula,\npróxima vacuna y control.", fill="#596770", font=font(30))
    rounded(d, (805, 650, 1320, 715), 18, "#18a59f")
    d.text((965, 668), "Guardar atención", fill="#ffffff", font=font(25, True))
    img.save(path)
    return path


def save_mockup_inventory() -> Path:
    path = ASSET_DIR / "mockup_inventory.png"
    img = Image.new("RGB", (1500, 900), "#f7fbfb")
    d = ImageDraw.Draw(img)
    d.text((95, 80), "Inventario para veterinaria", fill="#141f27", font=font(50, True))
    rounded(d, (90, 170, 1410, 790), 24, "#ffffff", "#d3e6e7", 2)
    headers = ["Producto", "Proveedor", "Lote", "Cantidad", "Estado"]
    xs = [130, 430, 720, 940, 1140]
    for x, h in zip(xs, headers):
        d.text((x, 225), h, fill="#0b5f69", font=font(24, True))
    rows = [
        ("Vacuna triple felina", "Biológicos Andinos", "VF-8891", "9", "Seguro"),
        ("Analgésico mascota", "Vet Pharma", "AM-2240", "3", "Stock bajo"),
        ("Suero fisiológico", "Clínicos SAS", "SF-1020", "0", "Crítico"),
        ("Kit curación", "SurgiLine", "KS-4477", "72", "Seguro"),
    ]
    for i, row in enumerate(rows):
        y = 295 + i * 100
        d.line((125, y - 25, 1345, y - 25), fill="#d3e6e7", width=2)
        for x, value in zip(xs, row):
            color = "#141f27"
            if value == "Seguro":
                color = "#3d9c5f"
            if value == "Stock bajo":
                color = "#e3ae3a"
            if value == "Crítico":
                color = "#d14a4a"
            d.text((x, y), value, fill=color, font=font(22, value in {"Seguro", "Stock bajo", "Crítico"}))
    d.text((130, 705), "Incluye lotes, vencimientos, costos, proveedor y cadena de frío.", fill="#596770", font=font(27))
    img.save(path)
    return path


def add_bg(slide, color=COLORS["white"]):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, x, y, w, h, text, size=24, bold=False, color=COLORS["ink"], align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Segoe UI"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if align:
        p.alignment = align
    return shape


def add_bullets(slide, x, y, w, h, items, size=21, color=COLORS["ink"]):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Segoe UI"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(8)
    return shape


def add_chip(slide, x, y, text, fill=COLORS["soft"], color=COLORS["teal_dark"]):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.15), Inches(0.42))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = COLORS["line"]
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Segoe UI"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.CENTER
    return shape


def add_title(slide, title, subtitle=""):
    add_textbox(slide, 0.65, 0.45, 11.8, 0.5, "VELMORAX | Flujo veterinaria", 12, True, COLORS["teal_dark"])
    add_textbox(slide, 0.65, 0.9, 8.2, 0.8, title, 34, True, COLORS["ink"])
    if subtitle:
        add_textbox(slide, 0.68, 1.65, 9.3, 0.55, subtitle, 18, False, COLORS["muted"])


def add_card(slide, x, y, w, h, title, body, accent=COLORS["teal"]):
    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    rect.fill.solid()
    rect.fill.fore_color.rgb = COLORS["white"]
    rect.line.color.rgb = COLORS["line"]
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.color.rgb = accent
    add_textbox(slide, x + 0.22, y + 0.18, w - 0.4, 0.34, title, 17, True, COLORS["ink"])
    add_textbox(slide, x + 0.22, y + 0.62, w - 0.38, h - 0.75, body, 14, False, COLORS["muted"])
    return rect


def add_image(slide, path, x, y, w):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))


def build_presentation() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    metrics = load_metrics()
    assets = {
        "login": save_mockup_login(),
        "dashboard": save_mockup_dashboard(),
        "agenda": save_mockup_agenda(),
        "clinical": save_mockup_clinical(),
        "inventory": save_mockup_inventory(),
    }

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1
    slide = prs.slides.add_slide(blank)
    add_bg(slide, COLORS["soft"])
    if LOGO_PATH.exists():
        slide.shapes.add_picture(str(LOGO_PATH), Inches(0.85), Inches(0.65), height=Inches(1.25))
    add_textbox(slide, 0.85, 2.0, 8.3, 0.8, "Velmorax", 54, True, COLORS["teal_dark"])
    add_textbox(slide, 0.9, 2.85, 8.6, 1.1, "Exposición del flujo de veterinaria", 36, True, COLORS["ink"])
    add_textbox(slide, 0.92, 4.0, 7.1, 0.6, "Proyecto final para presentar al SENA", 22, False, COLORS["muted"])
    add_chip(slide, 0.92, 5.05, "Sin código")
    add_chip(slide, 3.25, 5.05, "Enfoque visual")
    add_chip(slide, 5.58, 5.05, "Estado actual")
    add_image(slide, assets["login"], 8.45, 0.9, 4.2)

    # 2
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Problema que resuelve", "Las veterinarias pequeñas o medianas suelen manejar agenda, historia, vacunas e inventario en lugares separados.")
    add_card(slide, 0.8, 2.35, 3.7, 3.7, "Desorden operativo", "Citas, responsables, historias y pagos quedan dispersos en libretas, chats o archivos diferentes.", COLORS["red"])
    add_card(slide, 4.85, 2.35, 3.7, 3.7, "Riesgo clínico", "Se pueden perder datos clave: peso, vacunas, controles, alergias, fórmula o seguimiento.", COLORS["yellow"])
    add_card(slide, 8.9, 2.35, 3.7, 3.7, "Inventario sensible", "Vacunas e insumos requieren lote, vencimiento, proveedor y cadena de frío visibles.", COLORS["teal"])

    # 3
    slide = prs.slides.add_slide(blank)
    add_bg(slide, COLORS["soft"])
    add_title(slide, "Solución propuesta", "Velmorax concentra el trabajo veterinario en un flujo visual: ingresar, agendar, atender, registrar y controlar insumos.")
    add_bullets(slide, 0.9, 2.35, 5.2, 3.4, [
        "Ingreso por flujo: el usuario elige Veterinaria antes de entrar.",
        "Tablero con citas, pacientes visibles, inventario y alertas.",
        "Historia clínica adaptada a mascota y responsable.",
        "Seguimiento de vacunas, peso, controles y pagos.",
        "Inventario conectado a lotes, vencimientos y cadena de frío.",
    ], 22)
    add_image(slide, assets["dashboard"], 6.45, 2.05, 5.9)

    # 4
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Inicio de sesión", "La primera decisión visual es escoger el flujo. Para esta exposición se muestra Veterinaria.")
    add_image(slide, assets["login"], 0.75, 2.15, 6.4)
    add_bullets(slide, 7.55, 2.15, 4.85, 3.5, [
        "Tarjeta de Veterinaria marcada como flujo activo.",
        "Campos sencillos: correo, clave y recordar sesión.",
        "Recuperación de clave disponible desde la pantalla de acceso.",
        "El botón cambia a “Entrar a Veterinaria”.",
    ], 22)

    # 5
    slide = prs.slides.add_slide(blank)
    add_bg(slide, COLORS["soft"])
    add_title(slide, "Recorrido del usuario", "La demo puede presentarse como una secuencia clara de trabajo diario.")
    steps = [
        ("1. Login", "Elige Veterinaria y accede con rol clínico."),
        ("2. Agenda", "Consulta citas, confirma estados y registra nuevas atenciones."),
        ("3. Paciente", "Crea o revisa la ficha de mascota y responsable."),
        ("4. Atención", "Registra motivo, diagnóstico, plan, peso, vacunas y control."),
        ("5. Inventario", "Revisa insumos críticos, lote, vencimiento y cadena de frío."),
    ]
    x = 0.85
    for title, body in steps:
        add_card(slide, x, 2.35, 2.25, 3.25, title, body, COLORS["teal"])
        x += 2.48

    # 6
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Tablero principal", "El tablero funciona como centro de operaciones para ver lo urgente primero.")
    add_image(slide, assets["dashboard"], 0.8, 2.0, 6.55)
    add_bullets(slide, 7.75, 2.05, 4.6, 3.4, [
        "Citas del día y estado de la agenda.",
        "Cantidad de pacientes registrados.",
        "Alertas de inventario crítico o bajo.",
        "Cobertura por organización o sede.",
        "Actividad reciente para trazabilidad.",
    ], 22)

    # 7
    slide = prs.slides.add_slide(blank)
    add_bg(slide, COLORS["soft"])
    add_title(slide, "Agenda veterinaria", "Permite organizar el día de trabajo y reducir llamadas o citas perdidas.")
    add_image(slide, assets["agenda"], 0.75, 1.95, 6.5)
    add_card(slide, 7.65, 2.2, 4.7, 1.25, "Qué se ve", "Fecha, hora, mascota/responsable, servicio, sede, canal y estado de la cita.", COLORS["blue"])
    add_card(slide, 7.65, 3.75, 4.7, 1.25, "Qué permite", "Crear citas, actualizar estado y eliminar registros cuando el rol tiene permiso.", COLORS["teal"])
    add_card(slide, 7.65, 5.3, 4.7, 1.25, "Valor para la demo", "Muestra orden diario y priorización antes de abrir la atención clínica.", COLORS["green"])

    # 8
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Ficha de mascota", "El registro de paciente cambia cuando el tipo es Animal.")
    add_image(slide, assets["clinical"], 0.75, 1.95, 6.4)
    add_bullets(slide, 7.55, 2.05, 4.95, 3.8, [
        "Nombre visible de la mascota.",
        "Responsable o propietario.",
        "Especie, raza y peso.",
        "Estado de vacunación.",
        "Teléfono de contacto.",
        "Sede y organización asociada.",
    ], 22)

    # 9
    slide = prs.slides.add_slide(blank)
    add_bg(slide, COLORS["soft"])
    add_title(slide, "Atención clínica", "La historia veterinaria reúne información médica, seguimiento y recaudo.")
    add_card(slide, 0.85, 2.1, 3.8, 3.8, "Datos clínicos", "Motivo de consulta, nota, servicio realizado, diagnóstico, plan de tratamiento, alergias y signos vitales.", COLORS["teal"])
    add_card(slide, 4.95, 2.1, 3.8, 3.8, "Datos veterinarios", "Peso actual, próxima vacuna, fecha de control y estado de la atención: cerrada, seguimiento o abierta.", COLORS["green"])
    add_card(slide, 9.05, 2.1, 3.3, 3.8, "Cierre", "Fórmula, recomendaciones de salida, valor pagado, medio de pago y profesional.", COLORS["blue"])

    # 10
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Inventario de soporte", "El flujo veterinario se apoya en insumos controlados por lote, vencimiento y estado.")
    add_image(slide, assets["inventory"], 0.75, 1.95, 6.5)
    add_bullets(slide, 7.55, 2.05, 4.9, 3.6, [
        "Registro de vacunas, medicamentos e insumos.",
        "Proveedor, marca, código, lote y registro regulatorio.",
        "Semáforo: crítico, alerta o seguro.",
        "Cadena de frío para biológicos.",
        "Exportación de inventario para revisión externa.",
    ], 22)

    # 11
    slide = prs.slides.add_slide(blank)
    add_bg(slide, COLORS["soft"])
    add_title(slide, "Qué se lleva hasta el momento", "Estado actual según la base y las pantallas del proyecto.")
    metric_cards = [
        ("Pacientes", str(metrics["patients"])),
        ("Veterinaria", str(metrics["vet_patients"])),
        ("Historias", str(metrics["records"])),
        ("Atenciones vet", str(metrics["vet_records"])),
        ("Citas", str(metrics["appointments"])),
        ("Inventario", str(metrics["inventory"])),
    ]
    for i, (label, value) in enumerate(metric_cards):
        x = 0.9 + (i % 3) * 4.1
        y = 2.05 + (i // 3) * 1.7
        add_card(slide, x, y, 3.45, 1.25, label, value, COLORS["teal"])
    add_bullets(slide, 0.95, 5.65, 11.2, 0.8, [
        "También existen roles de usuario, recuperación de clave, auditoría, cierre diario, filtros por sede y compatibilidad con SQLite/PostgreSQL."
    ], 19, COLORS["muted"])

    # 12
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Qué faltaría por implementar", "Puntos recomendados para continuar después de la entrega visual.")
    add_card(slide, 0.85, 2.05, 3.7, 3.7, "Experiencia veterinaria", "Plantillas más completas por tipo de consulta, carné de vacunas imprimible y recordatorios automáticos.", COLORS["teal"])
    add_card(slide, 4.85, 2.05, 3.7, 3.7, "Operación", "Notificaciones por WhatsApp/SMS, adjuntos de imágenes, firma del profesional y reportes por periodo.", COLORS["blue"])
    add_card(slide, 8.85, 2.05, 3.7, 3.7, "Producción", "Despliegue final con dominio, HTTPS, PostgreSQL, respaldos programados y pruebas con usuarios reales.", COLORS["green"])

    # 13
    slide = prs.slides.add_slide(blank)
    add_bg(slide, COLORS["teal_dark"])
    add_textbox(slide, 1.0, 1.25, 10.8, 0.9, "Conclusión", 46, True, COLORS["white"])
    add_textbox(slide, 1.05, 2.25, 10.6, 1.65, "Velmorax ya presenta una base funcional para el flujo veterinario: ingreso dirigido, agenda, ficha de mascota, historia clínica, seguimiento e inventario.", 30, False, COLORS["white"])
    add_textbox(slide, 1.05, 4.45, 9.6, 0.85, "La siguiente etapa es fortalecer automatizaciones, reportes y despliegue para convertirlo en una herramienta lista para uso real.", 24, False, RGBColor(214, 239, 238))
    add_textbox(slide, 1.05, 6.35, 5.2, 0.35, "Proyecto final | SENA", 16, True, RGBColor(214, 239, 238))

    prs.save(PPTX_PATH)
    print(PPTX_PATH)


if __name__ == "__main__":
    build_presentation()
